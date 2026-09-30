# -*- coding: utf-8 -*-
"""Générateur hors ligne : classeur modifié (absence + affectation imposée), génération, exports, calendriers."""
import datetime as dt
import os
import zipfile
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from commun import Bilan, DIST, EXEMPLE, SORTIE, chromium, verifier_dist


def main():
    verifier_dist()
    B = Bilan("Générateur hors ligne")
    # classeur d'exemple + une absence (SB, 19-23/10/2026) + une affectation imposée (OA en IRM 2 le 12/10/2026)
    wb = load_workbook(EXEMPLE)
    ws = wb["Absences"]; ws["A5"], ws["B5"], ws["C5"], ws["D5"], ws["E5"] = "SB", dt.date(2026, 10, 19), dt.date(2026, 10, 23), "Journée", "Congés"
    ws = wb["Imposées"]; ws["A5"], ws["B5"], ws["C5"], ws["D5"] = dt.date(2026, 10, 12), "Matin", "OA", "IRM2"
    params = os.path.join(SORTIE, "params_test.xlsx"); wb.save(params)
    erreurs, reseau = [], []
    with sync_playwright() as p:
        b = chromium(p)
        ctx = b.new_context(accept_downloads=True, viewport={"width": 1280, "height": 900})
        ctx.route("**/*", lambda r: r.continue_() if r.request.url.startswith(("file:", "blob:", "data:"))
                  else (reseau.append(r.request.url), r.abort()))
        pg = ctx.new_page()
        pg.on("pageerror", lambda e: erreurs.append(str(e)))
        pg.goto("file://" + os.path.join(DIST, "generateur-hors-ligne.html"))
        pg.wait_for_selector("#srcinfo .chips")
        pg.set_input_files("#fileParams", params)
        pg.wait_for_function("document.querySelector('#srcinfo').innerText.includes('params_test')")
        B.ok("1 absence" in pg.inner_text("#srcinfo") and "imposée" in pg.inner_text("#srcinfo"), "classeur modifié lu")
        pg.select_option("#quality", "10"); pg.click("#btnGen")
        pg.wait_for_function("document.querySelector('#status').innerText.startsWith('Terminé')", timeout=90000)
        B.ok("erreur" not in pg.inner_text("#status"), "génération sans erreur")
        with pg.expect_download() as d:
            pg.click("#dlXlsx")
        xlsx = os.path.join(SORTIE, "hors_ligne.xlsx"); d.value.save_as(xlsx)
        with pg.expect_download() as d:
            pg.click("#dlZip")
        zp = os.path.join(SORTIE, "hors_ligne.zip"); d.value.save_as(zp)
        # recréation des calendriers depuis le planning produit
        pg.set_input_files("#filePlanning", xlsx)
        pg.wait_for_function("document.querySelector('#status2').innerText.includes('calendriers recréés')")
        B.ok(True, "calendriers recréés depuis un planning Excel")
        b.close()
    wb = load_workbook(xlsx)
    st, pp = wb["Statuts"], wb["Planning par poste"]
    col = next(c for c in range(4, st.max_column + 1) if st.cell(5, c).value == "SB")
    vals = [st.cell(r, col).value for r in range(7, st.max_row + 1)
            if isinstance(pp.cell(r, 1).value, dt.datetime) and 19 <= pp.cell(r, 1).value.day <= 23]
    B.ok(vals and set(vals) == {"ABS"}, "absence de SB respectée")
    irm2 = next(c for c in range(4, pp.max_column + 1) if pp.cell(6, c).value == "IRM2")
    ligne = next(r for r in range(7, pp.max_row + 1)
                 if isinstance(pp.cell(r, 1).value, dt.datetime) and pp.cell(r, 1).value.date() == dt.date(2026, 10, 12) and pp.cell(r, 3).value == "Matin")
    B.ok("OA" in str(pp.cell(ligne, irm2).value), "affectation imposée respectée (OA en IRM 2 le 12/10 matin)")
    z = zipfile.ZipFile(zp)
    B.ok(z.testzip() is None and len(z.namelist()) == 11, "archive des 11 calendriers valide")
    B.ok(not reseau, "aucune connexion réseau")
    B.ok(not erreurs, "aucune erreur JavaScript" + (f" : {erreurs[:3]}" if erreurs else ""))
    return B.fin()


if __name__ == "__main__":
    raise SystemExit(0 if main() else 1)
