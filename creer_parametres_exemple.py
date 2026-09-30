# -*- coding: utf-8 -*-
"""Construit un classeur de paramètres D'EXEMPLE (médecins fictifs).

Le vrai classeur du service ne doit jamais être déposé dans ce dépôt public.
Usage : python outils/creer_parametres_exemple.py [sortie.xlsx]
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

import os, sys
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "exemple", "parametres_exemple.xlsx")
os.makedirs(os.path.dirname(os.path.abspath(OUT)), exist_ok=True)

F = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(name=F, bold=True, color="FFFFFF", size=10)
IN_FILL = PatternFill("solid", fgColor="FFF2CC")   # cellules à saisir
T_FONT = Font(name=F, bold=True, size=14, color="1F3864")
N_FONT = Font(name=F, size=10)
I_FONT = Font(name=F, size=9, italic=True, color="7F7F7F")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

DEMI = [f"{j} {p}" for j in ["Lun", "Mar", "Mer", "Jeu", "Ven"] for p in ["M", "AM"]]


def header(ws, row, labels, widths=None):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=lab)
        c.fill, c.font, c.alignment, c.border = H_FILL, H_FONT, CENTER, BORDER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.row_dimensions[row].height = 32


def put(ws, row, values, inp=True, align=CENTER):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.border = N_FONT, BORDER
        c.alignment = align if i > 1 or align is CENTER else LEFT
        if inp:
            c.fill = IN_FILL


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = T_FONT
    ws["A2"] = sub
    ws["A2"].font = I_FONT


wb = Workbook()

# ---------------------------------------------------------------- Lisez-moi
ws = wb.active
ws.title = "Lisez-moi"
ws.column_dimensions["A"].width = 110
lines = [
    ("Planning du service de radiologie — classeur de paramètres", T_FONT),
    ("", N_FONT),
    ("Les cellules sur fond jaune sont à saisir ou à modifier. Le reste est calculé ou descriptif.", Font(name=F, bold=True, size=10)),
    ("", N_FONT),
    ("Onglets", Font(name=F, bold=True, size=11, color="1F3864")),
    ("• Paramètres : période à planifier, semaine de référence de l'alternance A/B, règles générales.", N_FONT),
    ("• Sites : horaires des vacations et adresse (reprise dans les calendriers iPhone).", N_FONT),
    ("• Postes : postes par site, nombre de médecins, demi-journées d'ouverture (X), priorité de couverture.", N_FONT),
    ("• Médecins : quotité, jours off (semaines A et B), indisponibilités fixes, compétences par poste.", N_FONT),
    ("• Absences : congés, formations, maladie… (une ligne par absence).", N_FONT),
    ("• Imposées : affectations à respecter obligatoirement (ex. : OA en IRM 2 le 12/10 matin).", N_FONT),
    ("• Fériés : jours non planifiés.", N_FONT),
    ("", N_FONT),
    ("Ajouter un médecin", Font(name=F, bold=True, size=11, color="1F3864")),
    ("Ajoutez une ligne dans l'onglet Médecins (initiales, quotité, off, compétences). Il sera pris en compte à la prochaine génération.", N_FONT),
    ("Pour un ajout ponctuel directement dans un planning déjà généré : voir l'onglet « Lisez-moi » du planning (colonnes libres).", N_FONT),
    ("", N_FONT),
    ("Priorités de couverture (1 = le plus important)", Font(name=F, bold=True, size=11, color="1F3864")),
    ("Quand il manque des médecins, les postes de priorité la plus élevée (chiffre le plus grand) ferment en premier.", N_FONT),
    ("La colonne « Priorité au-delà du minimum » s'applique aux places supplémentaires (ex. : 3e médecin au scanner).", N_FONT),
    ("", N_FONT),
    ("Compétences", Font(name=F, bold=True, size=11, color="1F3864")),
    ("Oui = peut tenir le poste · Préféré = affecté en priorité sur ce poste · vide ou Non = jamais affecté.", N_FONT),
    ("", N_FONT),
    ("Codes de demi-journée", Font(name=F, bold=True, size=11, color="1F3864")),
    ("M = matin, AM = après-midi. Indisponibilités fixes : « Jeu M » ; plusieurs : « Jeu M; Lun AM ».", N_FONT),
    ("Jours off : « Mercredi » ou « Mercredi, Jeudi ». Semaine A / B : alternance une semaine sur deux (voir Paramètres).", N_FONT),
]
for i, (t, f) in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=t).font = f
    ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")

# ---------------------------------------------------------------- Paramètres
ws = wb.create_sheet("Paramètres")
title(ws, "Paramètres généraux", "Modifiez les valeurs en jaune.")
header(ws, 4, ["Paramètre", "Valeur", "Commentaire"], [52, 18, 80])
params = [
    ("Début de période", dt.date(2026, 10, 1), "Premier jour du planning."),
    ("Fin de période", dt.date(2026, 10, 31), "Dernier jour du planning (les week-ends ne sont pas planifiés)."),
    ("Lundi de référence semaine A", dt.date(2026, 9, 28),
     "HYPOTHÈSE à vérifier : la semaine du 28/09/2026 est une semaine A (IA off mercredi, IB off vendredi)."),
    ("Un seul site par jour", "Oui", "Oui : un médecin qui part à Blaye ou Sainte-Foy y passe la journée (pas de trajet entre deux sites le même jour)."),
    ("Déplacements hors Libourne : cible par semaine (jours)", 1, "Chaque médecin va si possible une journée par semaine sur un autre site."),
    ("Déplacements hors Libourne : maximum par semaine (jours)", 2, "Au-delà, jamais."),
    ("Scanner interventionnel : minimum de vacations par semaine", 2, "Au moins 2 vacations tenues sur les 3 ouvertes."),
    ("Scanner interventionnel : référents (1 vacation chacun par semaine)", "IA, IB",
     "Chacun a 1 vacation par semaine ; un référent peut prendre les 2 si l'autre est absent."),
    ("Temps de calcul maximum (secondes)", 180, "Au-delà, le meilleur planning trouvé est retenu."),
]
for i, (k, v, c) in enumerate(params, 5):
    put(ws, i, [k, v, c], inp=False, align=LEFT)
    ws.cell(row=i, column=2).fill = IN_FILL
    ws.cell(row=i, column=2).alignment = CENTER
    if isinstance(v, dt.date):
        ws.cell(row=i, column=2).number_format = "DD/MM/YYYY"
dv = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=False)
ws.add_data_validation(dv)
dv.add("B8")

# ---------------------------------------------------------------- Sites
ws = wb.create_sheet("Sites")
title(ws, "Sites et horaires des vacations",
      "HYPOTHÈSE : horaires 8h30-13h00 / 13h30-18h00, à corriger. L'adresse apparaît dans les calendriers iPhone.")
header(ws, 4, ["Site", "Adresse (optionnel)", "Matin début", "Matin fin", "Après-midi début", "Après-midi fin"],
       [24, 50, 13, 13, 16, 16])
for i, s in enumerate(["Libourne", "Sainte-Foy-la-Grande", "Blaye"], 5):
    put(ws, i, [s, "", dt.time(8, 30), dt.time(13, 0), dt.time(13, 30), dt.time(18, 0)])
    for col in range(3, 7):
        ws.cell(row=i, column=col).number_format = "HH:MM"

# ---------------------------------------------------------------- Postes
ws = wb.create_sheet("Postes")
title(ws, "Postes par site",
      "X = poste ouvert sur cette demi-journée. Priorité : 1 = couvert en premier ; les chiffres élevés ferment en premier.")
cols = ["Code", "Libellé", "Site", "Groupe d'équilibrage", "Nb médecins souhaité", "Nb médecins minimum",
        "Priorité (jusqu'au minimum)", "Priorité au-delà du minimum", "Remarque"] + DEMI
header(ws, 4, cols, [9, 30, 20, 16, 11, 11, 12, 12, 42] + [7] * 10)
ALL = {d: "X" for d in DEMI}
def opn(days):
    return {d: ("X" if d in days else "") for d in DEMI}
postes = [
    ("SCAN", "Scanner (3 sites regroupés)", "Libourne", "Scanner", 3, 2, 1, 3,
     "3 médecins ; 2 seulement exceptionnellement.", ALL),
    ("SCI", "Scanner interventionnel", "Libourne", "Interventionnel", 1, 1, 8, 8,
     "Mar, Mer, Jeu matin. Minimum 2/semaine et 1 par référent : voir Paramètres.", opn(["Mar M", "Mer M", "Jeu M"])),
    ("SI", "Salle interventionnelle", "Libourne", "Interventionnel", 1, 1, 2, 2,
     "Fermée le mercredi. Jeudi matin : arthrographies (poste ARTH).",
     opn([d for d in DEMI if not d.startswith("Mer") and d != "Jeu M"])),
    ("ARTH", "Arthrographies (salle interv.)", "Libourne", "Interventionnel", 1, 1, 2, 2,
     "Jeudi matin, médecins ostéo-articulaires.", opn(["Jeu M"])),
    ("IRM1", "IRM 1", "Libourne", "IRM", 1, 1, 4, 4, "", ALL),
    ("IRM2", "IRM 2 (ostéo-articulaire)", "Libourne", "IRM ostéo", 1, 1, 7, 7,
     "Médecins ostéo-articulaires en priorité.", ALL),
    ("ECH1", "Échographie 1", "Libourne", "Échographie", 1, 1, 5, 5, "", ALL),
    ("ECH2", "Échographie 2", "Libourne", "Échographie", 1, 1, 9, 9, "Ferme en premier si manque de médecins.", ALL),
    ("MAM", "Mammographie Libourne", "Libourne", "Mammographie", 1, 1, 3, 3, "", ALL),
    ("SF-RE", "Radio-écho Sainte-Foy", "Sainte-Foy-la-Grande", "Échographie", 1, 1, 6, 6,
     "Fermeture seulement si vraiment pas assez de médecins.", ALL),
    ("BL-MAM", "Mammographie Blaye", "Blaye", "Mammographie", 1, 1, 6, 6,
     "Fermeture seulement si vraiment pas assez de médecins.", ALL),
    ("BL-RE", "Radio-écho Blaye", "Blaye", "Échographie", 1, 1, 6, 6,
     "Fermeture seulement si vraiment pas assez de médecins.", ALL),
]
for i, p in enumerate(postes, 5):
    vals = list(p[:9]) + [p[9][d] for d in DEMI]
    put(ws, i, vals)
    ws.cell(row=i, column=2).alignment = LEFT
    ws.cell(row=i, column=9).alignment = LEFT
    ws.cell(row=i, column=9).font = Font(name=F, size=9)
ws.freeze_panes = "C5"
dvx = DataValidation(type="list", formula1='"X"', allow_blank=True)
ws.add_data_validation(dvx)
dvx.add("J5:S60")
dvp = DataValidation(type="whole", operator="between", formula1="1", formula2="9")
ws.add_data_validation(dvp)
dvp.add("G5:H60")

# ---------------------------------------------------------------- Médecins
ws = wb.create_sheet("Médecins")
title(ws, "Médecins, quotités et compétences",
      "Une ligne par médecin. Compétences : Oui / Préféré / vide. Pour ajouter un médecin, remplissez la première ligne vide.")
codes = [p[0] for p in postes]
mcols = ["Initiales", "Nom (optionnel)", "Statut", "Quotité", "Off semaine A", "Off semaine B",
         "Indisponibilités fixes", "Surspécialités", "Actif"] + codes
header(ws, 4, mcols, [10, 18, 14, 9, 16, 16, 18, 30, 7] + [8] * len(codes))

GEN = {"SCAN": "Oui", "IRM1": "Oui", "IRM2": "Oui", "ECH1": "Oui", "ECH2": "Oui", "SF-RE": "Oui", "BL-RE": "Oui"}
def comp(**kw):
    d = dict(GEN)
    d.update(kw)
    return d
MAMMO = {"MAM": "Oui", "BL-MAM": "Oui"}
OSTEO = {"IRM2": "Préféré", "ARTH": "Oui", "SCI": "Oui"}
INTERV = {"SI": "Oui", "SCI": "Oui"}
docs = [
    ("DA", "", "PH", 1.0, "Lundi", "Lundi", "", "ORL (scanner, IRM)", comp(SCAN="Préféré", IRM1="Préféré")),
    ("DB", "", "PH", 0.8, "Mercredi, Jeudi", "Mercredi, Jeudi", "", "Digestif — travaille lun, mar, ven", comp(SCAN="Préféré", IRM1="Préféré")),
    ("IA", "", "PH", 1.0, "Mercredi", "Vendredi", "", "Radiologie interventionnelle, digestif", comp(**INTERV)),
    ("IB", "", "PH", 0.9, "Vendredi", "Mercredi", "Jeu M", "Radiologie interventionnelle, digestif — activité extérieure jeudi matin", comp(**INTERV)),
    ("SA", "", "Chef(fe) de service", 1.0, "Mercredi", "Mercredi", "", "Sénologie", comp(MAM="Préféré", **{"BL-MAM": "Oui"})),
    ("SB", "", "PH", 1.0, "Mercredi", "Mercredi", "", "Radiopédiatrie, sénologie", comp(ECH1="Préféré", ECH2="Préféré", **MAMMO)),
    ("SC", "", "Assistante", 1.0, "", "", "", "Radiopédiatrie, cardiaque, sénologie", comp(ECH1="Préféré", **MAMMO)),
    ("OA", "", "PH", 1.0, "Mardi", "Mardi", "", "Ostéo-articulaire", comp(**OSTEO)),
    ("OB", "", "PH", 1.0, "Jeudi", "Jeudi", "", "Ostéo-articulaire", comp(**OSTEO)),
    ("OC", "", "PH", 1.0, "Vendredi", "Vendredi", "", "Ostéo-articulaire", comp(**OSTEO)),
    ("SD", "", "PH", 1.0, "Lundi", "Lundi", "", "Sénologie", comp(MAM="Préféré", **{"BL-MAM": "Oui"})),
]
r = 5
for d in docs:
    vals = list(d[:8]) + ["Oui"] + [d[8].get(c, "") for c in codes]
    put(ws, r, vals)
    ws.cell(row=r, column=4).number_format = "0%"
    ws.cell(row=r, column=8).alignment = LEFT
    r += 1
# lignes vides prêtes pour de nouveaux médecins
for rr in range(r, r + 8):
    put(ws, rr, [""] * len(mcols))
ws.cell(row=11, column=5).comment = Comment(
    "Exemple : médecin sans jour off fixe.", "Planning")
ws.cell(row=6, column=5).comment = Comment(
    "Exemple : temps partiel à 80 %.", "Planning")
ws.freeze_panes = "B5"
dvc = DataValidation(type="list", formula1='"Oui,Préféré,Non"', allow_blank=True)
ws.add_data_validation(dvc)
dvc.add(f"J5:{ws.cell(row=60, column=len(mcols)).coordinate}")
dva = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
ws.add_data_validation(dva)
dva.add("I5:I60")

# ---------------------------------------------------------------- Absences
ws = wb.create_sheet("Absences")
title(ws, "Absences (congés, formation, maladie…)",
      "Exemple de saisie : SB | 19/10/2026 | 23/10/2026 | Journée | Congés. Période : Journée, Matin ou Après-midi.")
header(ws, 4, ["Initiales", "Du", "Au", "Période", "Motif"], [11, 14, 14, 14, 40])
for rr in range(5, 65):
    put(ws, rr, [""] * 5)
    ws.cell(row=rr, column=2).number_format = "DD/MM/YYYY"
    ws.cell(row=rr, column=3).number_format = "DD/MM/YYYY"
dvp2 = DataValidation(type="list", formula1='"Journée,Matin,Après-midi"', allow_blank=True)
ws.add_data_validation(dvp2)
dvp2.add("D5:D200")

# ---------------------------------------------------------------- Imposées
ws = wb.create_sheet("Imposées")
title(ws, "Affectations imposées",
      "Exemple de saisie : 12/10/2026 | Matin | OA | IRM2. Le code poste doit exister dans l'onglet Postes.")
header(ws, 4, ["Date", "Période", "Initiales", "Code poste", "Commentaire"], [14, 14, 11, 12, 40])
for rr in range(5, 45):
    put(ws, rr, [""] * 5)
    ws.cell(row=rr, column=1).number_format = "DD/MM/YYYY"
dvp3 = DataValidation(type="list", formula1='"Matin,Après-midi"', allow_blank=True)
ws.add_data_validation(dvp3)
dvp3.add("B5:B200")

# ---------------------------------------------------------------- Fériés
ws = wb.create_sheet("Fériés")
title(ws, "Jours fériés", "Jours non planifiés.")
header(ws, 4, ["Date", "Libellé"], [14, 30])
feries = [
    (dt.date(2026, 11, 1), "Toussaint"), (dt.date(2026, 11, 11), "Armistice"),
    (dt.date(2026, 12, 25), "Noël"), (dt.date(2027, 1, 1), "Jour de l'an"),
    (dt.date(2027, 3, 29), "Lundi de Pâques"), (dt.date(2027, 5, 1), "Fête du travail"),
    (dt.date(2027, 5, 6), "Ascension"), (dt.date(2027, 5, 8), "Victoire 1945"),
    (dt.date(2027, 5, 17), "Lundi de Pentecôte"), (dt.date(2027, 7, 14), "Fête nationale"),
    (dt.date(2027, 8, 15), "Assomption"),
]
for i, (d, l) in enumerate(feries, 5):
    put(ws, i, [d, l])
    ws.cell(row=i, column=1).number_format = "DD/MM/YYYY"

for s in wb.worksheets:
    s.sheet_view.showGridLines = False
wb.save(OUT)
print("OK", OUT)
