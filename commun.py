# -*- coding: utf-8 -*-
"""Outils communs aux tests (navigateur Chromium piloté par Playwright)."""
import calendar
import datetime as dt
import os
import sys
import tempfile

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DIST = os.path.join(RACINE, "dist")
EXEMPLE = os.path.join(RACINE, "exemple", "parametres_exemple.xlsx")
SORTIE = os.environ.get("TEST_SORTIE") or os.path.join(tempfile.gettempdir(), "planning-tests")
os.makedirs(SORTIE, exist_ok=True)


class Bilan:
    def __init__(self, nom):
        self.nom, self.echecs = nom, []
        print(f"\n=== {nom} ===")

    def ok(self, cond, libelle):
        print(("  OK     " if cond else "  ÉCHEC  ") + libelle)
        if not cond:
            self.echecs.append(libelle)

    def fin(self):
        print(f"--- {self.nom} : " + ("tout est OK" if not self.echecs else f"{len(self.echecs)} échec(s)"))
        return not self.echecs


def mois_suivant():
    """Premier et dernier jour du mois suivant : les tests restent valides quelle que soit la date."""
    t = dt.date.today()
    y, m = (t.year + 1, 1) if t.month == 12 else (t.year, t.month + 1)
    return dt.date(y, m, 1), dt.date(y, m, calendar.monthrange(y, m)[1])


def feries_exemple():
    from openpyxl import load_workbook
    ws = load_workbook(EXEMPLE, data_only=True)["Fériés"]
    out = set()
    for r in range(5, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, dt.datetime):
            out.add(v.date())
        elif isinstance(v, dt.date):
            out.add(v)
    return out


def jours_ouvres(d1, d2):
    fer = feries_exemple()
    n, d = [], d1
    while d <= d2:
        if d.weekday() < 5 and d not in fer:
            n.append(d)
        d += dt.timedelta(days=1)
    return n


def numero_jour(d):
    """Numéro de jour utilisé par l'application (jours depuis le 01/01/1970)."""
    return (d - dt.date(1970, 1, 1)).days


def chromium(p):
    chemin = os.environ.get("CHROMIUM_PATH")
    return p.chromium.launch(executable_path=chemin) if chemin else p.chromium.launch()


def verifier_dist():
    for f in ("index.html", "demo.html", "generateur-hors-ligne.html"):
        if not os.path.exists(os.path.join(DIST, f)):
            sys.exit(f"dist/{f} absent : lancez d'abord python outils/construire.py")
