# -*- coding: utf-8 -*-
"""
Règles du service (octobre 2026), sur un classeur d'essai dérivé de l'exemple (médecins fictifs) :
- ouverture d'un poste une semaine sur deux (A / B) ; poste « toujours en télétravail » ;
- même médecin le matin et l'après-midi : interdit / à éviter / journée entière ;
- vacations spécialisées (médecins habilités, minimum par semaine, condition « si présents ») ;
- préférences (poste ou site certains jours, demi-journée libre) ;
- notes récurrentes et astreintes (affichées ; médecin d'astreinte à l'IRM 1 l'après-midi) ;
- binôme jamais absent le même jour (repos déplacé au jour de repli).
Puis l'édition de ces onglets dans Admin → Règles et postes.
"""
import base64, datetime as dt, io, json, os
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from commun import Bilan, DIST, EXEMPLE, SORTIE, chromium

B = Bilan("Règles du service (vacations, préférences, notes, binôme)")
DEMI = [f"{j} {h}" for j in ("Lun", "Mar", "Mer", "Jeu", "Ven") for h in ("M", "AM")]

wb = load_workbook(EXEMPLE)
po = wb["Postes"]
hdr = [c.value for c in po[4]]
col = {h: i + 1 for i, h in enumerate(hdr)}
mj = col["Matin + après-midi"]
for r in range(5, po.max_row + 1):
    code = po.cell(r, 1).value
    if code == "IRM1": po.cell(r, mj, "Interdit")
    if code == "MAM": po.cell(r, mj, "Éviter")
    if code == "SI": po.cell(r, mj, "Journée entière")
r = po.max_row + 1
ouv = {"Lun M": "B", "Lun AM": "B", "Mar M": "A", "Mar AM": "A", "Jeu M": "X", "Jeu AM": "X", "Ven M": "A", "Ven AM": "A"}
for k, v in {"Code": "IRMB", "Libellé": "IRM Blaye (interprétation)", "Site": "Libourne", "Groupe d'équilibrage": "IRM",
             "Nb médecins souhaité": 1, "Nb médecins minimum": 1, "Priorité (jusqu'au minimum)": 5, "Priorité au-delà du minimum": 5,
             "Télétravail possible": "Toujours", **ouv}.items():
    po.cell(r, col[k], v)
po.cell(r, mj, "Journée entière")
me = wb["Médecins"]
c = me.max_column + 1
me.cell(4, c, "IRMB")
for rr in range(5, me.max_row + 1):
    if me.cell(rr, 1).value: me.cell(rr, c, "Oui")
pa = wb["Paramètres"]
rr = pa.max_row + 1
if not any(pa.cell(k, 1).value == "Binôme : jamais absents le même jour" for k in range(5, rr)):
    pa.cell(rr, 1, "Binôme : jamais absents le même jour"); pa.cell(rr, 2, "IA, IB")
    pa.cell(rr + 1, 1, "Binôme : jour de repos de repli"); pa.cell(rr + 1, 2, "Mercredi")
ab = wb["Absences"]
for i, v in enumerate(["IB", "16/11/2026", "20/11/2026", "Journée", "Congés"], 1): ab.cell(5, i, v)
def onglet(nom, entetes, lignes):
    ws = wb[nom] if nom in wb.sheetnames else wb.create_sheet(nom)
    ws.cell(1, 1, nom)
    for i, h in enumerate(entetes, 1): ws.cell(4, i, h)
    for k, l in enumerate(lignes, 5):
        for i, v in enumerate(l, 1): ws.cell(k, i, v)
onglet("Vacations spécialisées", ["Nom", "Postes", "Créneaux", "Médecins habilités", "Minimum par semaine", "Si présents", "Remarque"],
       [["Coro", "SCAN", "Lun AM, Ven M", "SC, DA", 1, "", ""], ["Pédiatrie IRM", "IRM1", "Jeu M", "SB, SC", None, "", ""],
        ["Pédiatrie écho supplémentaire", "ECH1, ECH2", "Mer M", "SB, SC", 0, "SB, SC", "SB est off le mercredi : jamais ouverte"]])
onglet("Préférences", ["Médecin", "Type", "Poste ou site", "Jours", "Remarque"],
       [["OA", "Poste", "Blaye", "Lundi", ""], ["SA", "Libre", "", "Mar M, Mar AM", "Bureau"]])
onglet("Notes", ["Texte", "Quand", "Demi-journée"], [["Réunion de service 13h30", "1er mardi", ""], ["Staff scanner", "3e jeudi", "Après-midi"]])
onglet("Astreintes", ["Date", "Médecin"], [["03/11/2026", "DA"], ["10/11/2026", "DA"], ["17/11/2026", "DA"], ["24/11/2026", "DA"]])
essai = os.path.join(SORTIE, "parametres_regles_service.xlsx")
wb.save(essai)
b64 = base64.b64encode(open(essai, "rb").read()).decode()

CALCUL = """async (b64) => {
  const bytes = b64ToBytes(b64);
  const P = readParams(await readXlsx(bytes.buffer));
  P.start = dayFromYMD(2026, 11, 2); P.end = dayFromYMD(2026, 11, 27);
  const pr = buildProblem(P);
  const R = await solve(pr, 12);
  const items = checks(pr, R);
  const draft = draftFromResult(pr, R);
  const xl = await exportPlanning(pr, R, items, {});
  const get = (s, code) => R.assign.get(s + "|" + pr.pIdx[code]) || [];
  const jour = (d) => pr.slots.map((x, i) => i).filter((i) => pr.slots[i].d === d);
  const out = { ouvertIRMB: {}, irmbJours: 0, irmbMemeMedecin: 0, irmbSurSite: 0, irm1Double: 0, siJours: 0, siMeme: 0,
    pediaIRM: [], coroSemaines: [], prefOA: [], libreSA: [], astreinte: [], items, notes: draft.notes,
    vacs: pr.vacs.map((v) => v.nom), xl: bytesToB64(xl) };
  for (const d of pr.days) {
    const [m, a] = jour(d), wt = pr.weekType.get(mondayOf(d)), wd = weekday(d);
    out.ouvertIRMB[wt + wd] = pr.open[m][pr.pIdx.IRMB];
    if (pr.open[m][pr.pIdx.IRMB]) {
      const xm = get(m, "IRMB-TT"), xa = get(a, "IRMB-TT");
      out.irmbSurSite += get(m, "IRMB").length + get(a, "IRMB").length;
      if (xm.length || xa.length) out.irmbJours++;
      if (xm.length && xm[0] === xa[0]) out.irmbMemeMedecin++;
    }
    for (const ini of get(m, "IRM1")) if (get(a, "IRM1").includes(ini)) out.irm1Double++;
    if (pr.open[m][pr.pIdx.SI] && pr.open[a][pr.pIdx.SI] && get(m, "SI").length) { out.siJours++; if (get(m, "SI")[0] === get(a, "SI")[0]) out.siMeme++; }
    if (wd === 3) out.pediaIRM.push(get(m, "IRM1").some((x) => ["SB", "SC"].includes(x)));
  }
  pr.weeks.forEach((w, wi) => {
    const ss = pr.slots.map((x, i) => i).filter((i) => pr.slots[i].w === wi);
    out.coroSemaines.push(ss.filter((i) => (weekday(pr.slots[i].d) === 0 && pr.slots[i].h === "AM") || (weekday(pr.slots[i].d) === 4 && pr.slots[i].h === "M"))
      .some((i) => get(i, "SCAN").some((x) => ["SC", "DA"].includes(x))));
    const lun = ss.filter((i) => weekday(pr.slots[i].d) === 0), oa = pr.docs.indexOf("OA");
    if (lun.some((i) => pr.status[oa][i] === "")) out.prefOA.push(lun.some((i) => ["BL-MAM", "BL-RE"].some((c) => get(i, c).includes("OA"))));
    const mar = ss.filter((i) => weekday(pr.slots[i].d) === 1), sa = pr.docs.indexOf("SA");
    if (mar.length) out.libreSA.push(mar.some((i) => !pr.postes.some((c) => get(i, c).includes("SA"))));
  });
  for (const d of [dayFromYMD(2026, 11, 3), dayFromYMD(2026, 11, 10), dayFromYMD(2026, 11, 17), dayFromYMD(2026, 11, 24)])
    out.astreinte.push(get(jour(d)[1], "IRM1").includes("DA"));
  const ia = pr.docs.indexOf("IA");
  out.binome = { mer: jour(dayFromYMD(2026, 11, 18)).map((i) => pr.status[ia][i]), ven: jour(dayFromYMD(2026, 11, 20)).map((i) => pr.status[ia][i]) };
  return out;
}"""

with sync_playwright() as p:
    b = chromium(p)
    ctx = b.new_context(viewport={"width": 390, "height": 844}, is_mobile=True, has_touch=True)
    pg = ctx.new_page(); errs = []
    pg.on("pageerror", lambda e: errs.append(str(e))); pg.on("dialog", lambda d: d.accept())
    pg.goto("file://" + os.path.join(DIST, "demo.html") + "#admin")
    pg.wait_for_selector("#loginCard:not([hidden])")
    pg.fill("#loginEmail", "a@b.fr"); pg.fill("#loginPw", "x"); pg.click("#loginForm button[type=submit]")
    pg.wait_for_selector("#adminBody:not([hidden])")
    r = pg.evaluate(CALCUL, b64)

    o = r["ouvertIRMB"]
    B.ok(o.get("A1") and o.get("A3") and o.get("A4") and not o.get("A0") and o.get("B0") and o.get("B3") and not o.get("B1") and not o.get("B4"),
         f"ouverture A / B : IRM Blaye mar-jeu-ven en semaine A, lun-jeu en semaine B ({o})")
    B.ok(r["irmbSurSite"] == 0, "poste « toujours en télétravail » : jamais tenu sur site")
    B.ok(r["irmbMemeMedecin"] >= r["irmbJours"] - 1, f"journée entière : même médecin matin et après-midi ({r['irmbMemeMedecin']}/{r['irmbJours']} jours tenus)")
    B.ok(r["siMeme"] >= r["siJours"] - 2, f"salle interventionnelle à la journée ({r['siMeme']}/{r['siJours']})")
    B.ok(r["irm1Double"] == 0, "IRM 1 : jamais le même médecin matin et après-midi (interdit)")
    B.ok(all(r["pediaIRM"]), f"vacation obligatoire : pédiatrie à l'IRM chaque jeudi matin ({r['pediaIRM']})")
    B.ok(all(r["coroSemaines"]), f"vacation « 1 par semaine » : coro au scanner assurée chaque semaine ({r['coroSemaines']})")
    B.ok("Pédiatrie écho supplémentaire" not in r["vacs"], "condition « si présents » : vacation absente quand un médecin manque")
    B.ok(sum(r["prefOA"]) >= len(r["prefOA"]) - 1 and r["prefOA"], f"préférence : OA à Blaye le lundi ({r['prefOA']})")
    B.ok(all(r["libreSA"]), f"préférence : SA garde une demi-journée libre le mardi ({r['libreSA']})")
    B.ok(sum(r["astreinte"]) >= 3, f"astreinte : DA à l'IRM 1 l'après-midi ({r['astreinte']})")
    B.ok(r["binome"]["mer"] == ["OFF", "OFF"] and r["binome"]["ven"] == ["", ""], f"binôme : IB absent, repos de IA déplacé du vendredi au mercredi ({r['binome']})")
    infos = [t for niv, t in r["items"] if niv == "Info"]
    B.ok(any("repos déplacé" in t for t in infos), "binôme : déplacement signalé dans les contrôles")
    erreurs = [t for niv, t in r["items"] if niv == "Erreur"]
    B.ok(not erreurs, f"contrôles : aucune erreur {erreurs[:3]}")
    jn = lambda y, m, d: str((dt.date(y, m, d) - dt.date(1970, 1, 1)).days)
    n = r["notes"] or {}
    B.ok("Réunion de service 13h30" in n.get(jn(2026, 11, 3), []) and "Staff scanner (après-midi)" in n.get(jn(2026, 11, 19), [])
         and "Astreinte : DA" in n.get(jn(2026, 11, 10), []), "notes : 1er mardi, 3e jeudi et astreintes dans le planning publié")
    x = load_workbook(io.BytesIO(base64.b64decode(r["xl"])))["Planning par poste"]
    B.ok(any("Réunion de service" in str(c.value or "") for row in x.iter_rows() for c in row), "notes : visibles dans l'Excel")

    # --- édition dans le site
    pg.set_input_files("#upParams", essai)
    pg.wait_for_selector("#regList [data-reg]")
    txt = pg.inner_text("#regList")
    B.ok(all(k in txt for k in ("Vacations spécialisées", "Préférences", "Notes du planning", "Astreintes")), "Règles et postes : nouveaux onglets listés")
    pg.click('#regList [data-reg="Postes"]'); pg.wait_for_selector("#regOverlay:not([hidden])")
    pg.locator("#regItems li", has_text="IRM2").locator("[data-edit]").click()
    pg.select_option("select[data-mode-for]", "B")
    pg.click('#regForm .chips.cycle button[data-k="0"]')          # Lun M : X -> semaine B
    pg.click('#regForm .chips.cycle button[data-k="1"]')          # Lun AM : X -> semaine B
    pg.select_option("select[data-mode-for]", "X")
    pg.click('#regForm .chips.cycle button[data-k="9"]')          # Ven AM : X -> fermé
    lab = pg.inner_text('#regForm .chips.cycle button[data-k="0"]')
    pg.select_option("#rf10", "Interdit")
    pg.click("#regForm button[type=submit]")
    pg.wait_for_function("document.querySelector('#regFormMsg').innerText.startsWith('Enregistré')", timeout=10000)
    pg.wait_for_timeout(300)
    q = pg.evaluate("""async () => { const P = readParams(await readXlsx(b64ToBytes(S.params.xlsx).buffer)); const x = P.postes.IRM2;
        return { opens: [...x.opens], week: Object.fromEntries(x.opensWeek), mj: x.journee }; }""")
    B.ok("Lun M · B" == lab and q["week"] == {"0M": "B", "0AM": "B"} and "4AM" not in q["opens"] and q["mj"] == "demi",
         f"éditeur de postes : semaine B seulement, fermeture, « Interdit » enregistrés ({q})")
    pg.click("#regClose")
    pg.click('#regList [data-reg="Vacations spécialisées"]'); pg.wait_for_selector("#regOverlay:not([hidden])")
    pg.click("#regAdd")
    pg.fill("#rf0", "IRM cardiaque")
    pg.click('#rf1 button[data-k="IRM1"]'); pg.click('#rf2 button[data-k="7"]')      # Jeu AM
    pg.click('#rf3 button[data-k="SC"]')
    pg.click("#regForm button[type=submit]")
    pg.wait_for_function("document.querySelector('#regFormMsg').innerText.startsWith('Enregistré')", timeout=10000)
    pg.wait_for_timeout(300)
    v = pg.evaluate("""async () => { const P = readParams(await readXlsx(b64ToBytes(S.params.xlsx).buffer));
        const v = P.vacs.find((x) => x.nom === 'IRM cardiaque'); return v && { postes: v.postes, cr: [...v.creneaux], hab: v.hab, min: v.min }; }""")
    B.ok(v == {"postes": ["IRM1"], "cr": ["3AM"], "hab": ["SC"], "min": None}, f"éditeur : vacation spécialisée ajoutée ({v})")
    pg.click("#regClose")
    pg.click('#regList [data-reg="Notes"]'); pg.wait_for_selector("#regOverlay:not([hidden])")
    pg.click("#regAdd"); pg.fill("#rf0", "Staff IRM"); pg.fill("#rf1", "1er jeudi")
    pg.click("#regForm button[type=submit]")
    pg.wait_for_function("document.querySelector('#regFormMsg').innerText.startsWith('Enregistré')", timeout=10000)
    B.ok("Staff IRM" in pg.inner_text("#regItems"), "éditeur : note ajoutée")
    pg.click("#regClose")
    B.ok(not errs, "aucune erreur JavaScript" + (f" : {errs[:3]}" if errs else ""))
    b.close()
raise SystemExit(0 if B.fin() else 1)
