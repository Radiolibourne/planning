# -*- coding: utf-8 -*-
"""
Générateur hors ligne : astreintes. Indisponibilités saisies sur le calendrier, enregistrées dans l'onglet
« Indisponibilités astreinte » du classeur ; génération ; astreintes enregistrées dans l'onglet « Astreintes ».
"""
import os
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from commun import Bilan, DIST, SORTIE, chromium

B = Bilan("Générateur hors ligne : astreintes")
with sync_playwright() as p:
    b = chromium(p)
    pg = b.new_context(viewport={"width": 1200, "height": 900}, accept_downloads=True).new_page(); errs = []
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto("file://" + os.path.join(DIST, "generateur-hors-ligne.html"))
    pg.wait_for_selector("#aTab tbody tr", timeout=15000)
    B.ok(pg.locator("#aTab tbody tr").count() >= 10, "médecins du classeur listés")
    pg.fill("#aDu", "2027-01-01"); pg.fill("#aAu", "2027-03-31"); pg.dispatch_event("#aAu", "change")
    pg.fill("#aNouv", "Rayann"); pg.click("#aAjout")
    i = pg.evaluate("AST.lignes.findIndex((l) => l.id === 'DA')")
    pg.click(f'#aTab [data-cal="{i}"]'); pg.wait_for_selector("#aCal .cj")
    jours = pg.evaluate("[dayFromYMD(2027, 1, 8), dayFromYMD(2027, 1, 9), dayFromYMD(2027, 2, 3)]")
    for d in jours: pg.click(f'#aCal .cj[data-j="{d}"]')
    B.ok("3 jour(s)" in pg.locator("#aTab tbody tr").nth(i).inner_text(), "calendrier : jours d'indisponibilité comptés")
    with pg.expect_download() as dl:
        pg.click("#aSaveWb")
    f1 = os.path.join(SORTIE, "hl_indispos.xlsx"); dl.value.save_as(f1)
    ws = load_workbook(f1)["Indisponibilités astreinte"]
    lignes = {r[0]: r for r in ws.iter_rows(min_row=5, values_only=True) if r[0]}
    B.ok(lignes["DA"][3] == "08/01/2027-09/01/2027 ; 03/02/2027" and "RAYANN" in lignes, f"onglet « Indisponibilités astreinte » écrit ({lignes['DA']})")
    pg.click("#aGen"); pg.wait_for_selector("#aSaveAst", timeout=15000)
    r = pg.evaluate("[controlerAstreintes(AST.res.res, AST.res.indispo, 3), AST.res.res.aPourvoir.length, AST.res.res.unites.filter((u) => u.qui === 'DA').flatMap((u) => u.jours)]")
    B.ok(not r[0] and r[1] == 0 and not set(r[2]) & set(jours), "génération : règles et indisponibilités respectées")
    with pg.expect_download() as dl:
        pg.click("#aSaveAst")
    f2 = os.path.join(SORTIE, "hl_astreintes.xlsx"); dl.value.save_as(f2)
    wb = load_workbook(f2)
    n = sum(1 for row in wb["Astreintes"].iter_rows(min_row=5, values_only=True) if row[0])
    B.ok(n == 90 and "Indisponibilités astreinte" in wb.sheetnames, f"astreintes enregistrées dans le classeur, un jour par ligne ({n})")
    B.ok(pg.evaluate("current.P.astreintes.length") == 90, "le classeur rechargé : la génération du planning tient compte des astreintes")
    B.ok(not errs, "aucune erreur JavaScript" + (f" : {errs[:3]}" if errs else ""))
    b.close()
raise SystemExit(0 if B.fin() else 1)
