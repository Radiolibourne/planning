# -*- coding: utf-8 -*-
"""
RCP : import des fiches PDF (scans avec texte reconnu), avec des médecins fictifs. Bandeau RCP dans Mon planning,
vue RCP dans Général, onglet Astreintes. La fiche est fabriquée ici à la manière d'un scanner : texte placé mot à mot,
dates et initiales abîmées par la reconnaissance (« 0z/ll/2026 », « DAIOA »), un flux compressé.
"""
import os, zlib
from openpyxl import Workbook
from playwright.sync_api import sync_playwright
from commun import Bilan, DIST, SORTIE, chromium

B = Bilan("RCP et astreintes")


def fiche_pdf(chemin, mots, compresse=True):
    """PDF d'une page : mots [(x, y, taille, texte)] en texte invisible (comme un scan avec reconnaissance)."""
    lignes = ["BT", "3 Tr"]
    for x, y, s, t in mots:
        lignes += [f"/F3 {s} Tf", f"1 0 0 1 {x:.2f} {y:.2f} Tm", f"<{t.encode('cp1252').hex().upper()}> Tj"]
    lignes.append("ET")
    flux = "\n".join(lignes).encode("latin-1")
    corps = zlib.compress(flux) if compresse else flux
    filtre = " /Filter /FlateDecode" if compresse else ""
    objets = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F3 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        f"<< /Length {len(corps)}{filtre} >>\nstream\n".encode() + corps + b"\nendstream",
    ]
    out = b"%PDF-1.3\n"
    for i, o in enumerate(objets, 1):
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    out += b"trailer\n<< /Root 1 0 R >>\n%%EOF\n"
    with open(chemin, "wb") as f:
        f.write(out)


# Fiche de novembre 2026 : PNEUMO (16H-17H), SENO (LUNDI 13H-14H), URO (17H-20H 1er lundi / 16H-17H 3ème lundi)
mots = [(54, 777, 9.5, "Fait le 18 octobre 2026"), (221, 738, 14.5, "RCP "), (256, 738, 15.5, "NOVEMBRE "), (339, 737, 17.5, "2026"),
        (62, 621, 8, "SERVICE"), (150, 620, 10, "PNEUMO"), (280, 619, 9.5, "SENO"), (420, 618, 9.5, "URO"),
        (148, 563, 10, "t6H "), (170, 563, 10.5, "- "), (177, 563, 10, "17H"),
        (262, 562, 10, "TUNDI 13H-14H"),
        (400, 576, 11, "t7H - 20H"), (395, 561, 10.5, "1er lundi 16H."), (396, 545, 10, "17H 3ème lundi"),
        (52, 540, 8.5, "DATE"),
        (55, 500, 10, "0z/ll/2026"), (140, 515, 10, "Mme ALPHA"), (150, 503, 10, "Anne"), (145, 488, 10, "Poste 1234"), (275, 500, 10, "DAIOA"),
        (412, 515, 10, "Mr BETA Paul"), (418, 488, 10, "Poste 7270"),
        (55, 440, 10, "16/LL/2026"), (412, 455, 10, "MTBETA Paul"), (418, 430, 10, "Poste 7270"),
        (55, 380, 10, "23/77/2026"), (275, 380, 10, "SB/]B"),
        (120, 300, 11, "Si vous n'arrivez pas à joindre le radiologue merci de bien vouloir appeler au 3129")]
FICHE = os.path.join(SORTIE, "rcp_novembre.pdf")
fiche_pdf(FICHE, mots)
FICHE2 = os.path.join(SORTIE, "rcp_novembre_bis.pdf")
fiche_pdf(FICHE2, [m for m in mots if m[3] not in ("Mr BETA Paul", "Poste 7270", "MTBETA Paul")], compresse=False)

# planning publié de novembre (fait à la main), avec astreintes
wb = Workbook(); wb.remove(wb.active)
LIB = {7: "DEPLACEMENT BLAYE", 10: "SCANNER : LIBOURNE, BLAYE, STE FOY", 15: "IRM 1 + RADIO URGENCES", 32: "ASTREINTE"}
for titre, entete, cases in [("DU 2 NOVEMBRE AU 6 NOVEMBRE", "SEMAINE DU 2 NOVEMBRE 2026 AU 6 NOVEMBRE 2026", {(10, 2): "DA/OA", (15, 2): "IA", (32, 2): "IA", (32, 4): "OA", (32, 6): "IA"}),
                             ("DU 16 NOVEMBRE AU 20 NOVEMBRE", "SEMAINE DU 16 NOVEMBRE 2026 AU 20 NOVEMBRE 2026", {(10, 2): "DB/OB", (32, 2): "DA"}),
                             ("DU 23 NOVEMBRE AU 27 NOVEMBRE", "SEMAINE DU 23 NOVEMBRE 2026 AU 27 NOVEMBRE 2026", {(10, 2): "SB/OC", (32, 2): "Rayann", (32, 4): "RAYANN"})]:
    ws = wb.create_sheet(titre); ws["A1"] = entete
    for r, t in LIB.items(): ws.cell(r, 1, t)
    for (r, c), v in cases.items(): ws.cell(r, c, v)
PLANNING = os.path.join(SORTIE, "planning_rcp.xlsx"); wb.save(PLANNING)

with sync_playwright() as p:
    b = chromium(p)
    pg = b.new_context(viewport={"width": 390, "height": 844}, is_mobile=True, has_touch=True).new_page(); errs = []
    pg.on("pageerror", lambda e: errs.append(str(e))); pg.on("dialog", lambda d: d.accept())
    pg.add_init_script("Date.now = (() => { const f = Date.now; const off = new Date('2026-11-02T09:00:00Z').getTime() - f(); return () => f() + off; })();")
    pg.goto("file://" + os.path.join(DIST, "demo.html") + "#admin")
    pg.wait_for_selector("#loginCard:not([hidden])")
    pg.fill("#loginEmail", "a@b.fr"); pg.fill("#loginPw", "x"); pg.click("#loginForm button[type=submit]")
    pg.wait_for_selector("#adminBody:not([hidden])")
    pg.click("#useDefault"); pg.wait_for_function("document.querySelector('#paramInfo').innerText.includes('médecins')")
    pg.set_input_files("#upPlanning", PLANNING)
    pg.wait_for_function("document.querySelector('#gMsg').innerText.startsWith('Planning importé')", timeout=10000)
    pg.click("#publishBtn"); pg.wait_for_function("S.published && S.published.jours.length > 0", timeout=10000)
    # import de la fiche : lecture, noms à associer, aperçu
    pg.set_input_files("#upRcp", FICHE)
    pg.wait_for_selector("#rcpSave", timeout=10000)
    lu = pg.evaluate("""() => { const l = RCP_ATTENTE.lots[0]; return { mois: l.mois, an: l.an, cols: l.colonnes, inconnus: RCP_ATTENTE.inconnus,
        items: l.items.map((i) => [i.jour, l.colonnes[i.col].service, i.ini.join('/'), i.inconnus.join(',')]) }; }""")
    B.ok(lu["mois"] == 11 and lu["an"] == 2026 and [c["service"] for c in lu["cols"]] == ["PNEUMO", "SENO", "URO"], f"fiche lue : mois et colonnes ({lu['cols']})")
    B.ok(lu["cols"][0]["horaire"] == "16h-17h" and lu["cols"][1]["horaire"] == "LUNDI 13h-14h" and "3ème lundi" in lu["cols"][2]["horaire"], "horaires lus (« t6H » → 16h)")
    B.ok(sorted(lu["inconnus"]) == ["ALPHA", "BETA"], f"noms à associer : {lu['inconnus']}")
    it = {(i[0], i[1]): i for i in lu["items"]}
    B.ok(set(it) == {(2, "PNEUMO"), (2, "SENO"), (2, "URO"), (16, "URO"), (23, "SENO")}, f"dates abîmées relues (« 0z/ll » → 2, « 23/77 » → 23) : {sorted(it)}")
    B.ok(it[(2, "SENO")][2] == "DA/OA" and it[(23, "SENO")][2] == "SB/JB" and it[(16, "URO")][3] == "BETA", f"initiales collées ou mal lues découpées ({it[(2, 'SENO')]}, {it[(23, 'SENO')]})")
    B.ok("Si vous" not in str(lu["items"]), "texte hors tableau ignoré")
    pg.select_option('[data-rcp-alias="ALPHA"]', "IA"); pg.select_option('[data-rcp-alias="BETA"]', "OB")
    B.ok("ALPHA ?" not in pg.inner_text("#rcpApercu") and pg.locator("#rcpApercu tbody tr").count() == 5, "aperçu : noms associés, 5 RCP")
    pg.click("#rcpSave"); pg.wait_for_function("S.rcp && S.rcp.items && S.rcp.items.length === 5", timeout=8000)
    rc = pg.evaluate("({ al: S.rcp.alias, it: S.rcp.items.map((x) => [ymd(x.d).d, x.service, x.ini.join('/')]) })")
    B.ok(rc["al"] == {"ALPHA": "IA", "BETA": "OB"} and [2, "PNEUMO", "IA"] in rc["it"] and [16, "URO", "OB"] in rc["it"], "RCP enregistrées, noms retenus")
    # 2e import (autre version de la fiche, sans le radiologue d'uro) : noms déjà connus, mêmes colonnes remplacées
    pg.set_input_files("#upRcp", FICHE2); pg.wait_for_selector("#rcpSave", timeout=10000)
    B.ok(pg.evaluate("RCP_ATTENTE.inconnus.length") == 0, "2e import : noms reconnus d'après les imports précédents")
    pg.click("#rcpSave"); pg.wait_for_function("S.rcp.items.length === 3", timeout=8000)
    B.ok("novembre 2026 · 3 RCP" in pg.inner_text("#rcpListe"), "ré-import : RCP des mêmes colonnes remplacées")
    # Mon planning : bandeau RCP
    pg.evaluate("S.me = 'IA'; location.hash = '#mon'"); pg.wait_for_selector("#vMon .rcpb", timeout=8000)
    t = pg.inner_text("#vMon .rcpb")
    B.ok("RCP" in t and "Pneumo" in t and "16h-17h" in t, f"Mon planning : bandeau RCP ({t})")
    pg.evaluate("S.me = 'OA'; renderMon()")
    B.ok("Séno" in pg.inner_text("#vMon .rcpb") and "13h-14h" in pg.inner_text("#vMon .rcpb") and "LUNDI" not in pg.inner_text("#vMon .rcpb"), "bandeau : horaire sans le jour")
    # Général → RCP
    pg.evaluate("location.hash = '#general'"); pg.wait_for_selector('#genBody .seg button[data-mode="rcp"]')
    pg.click('#genBody .seg button[data-mode="rcp"]'); pg.wait_for_selector("#genBody table.rcpt")
    g = pg.inner_text("#genBody table.rcpt")
    B.ok(g.count("Lun 02/11") == 1 and "Séno" in g and "DA / OA" in g and pg.locator("#genBody table.rcpt .me").count() == 1, "Général → RCP : tableau du mois, ma RCP en évidence")
    B.ok(pg.evaluate("document.documentElement.scrollWidth <= window.innerWidth"), "Général → RCP : pas de défilement horizontal sur téléphone")
    # horaire selon le rang du lundi (1er lundi / 3ème lundi)
    h = pg.evaluate("[horaireRcp('17h-20h 1er lundi 16h-17h 3ème lundi', dayFromYMD(2026, 11, 2)), horaireRcp('17h-20h 1er lundi 16h-17h 3ème lundi', dayFromYMD(2026, 11, 16))]")
    B.ok(h == ["17h-20h", "16h-17h"], f"horaire selon le 1er ou le 3e lundi ({h})")
    # onglet Astreintes
    pg.click("a[data-tab=astreintes]"); pg.wait_for_selector("#astBody table.astt", timeout=8000)
    a = pg.inner_text("#astBody")
    B.ok("Lun 02/11" in a and "Mer 04/11" in a and "Lun 23/11" in a and pg.locator("#astBody tbody tr").count() == 6, f"Astreintes : liste complète du mois ({pg.locator('#astBody tbody tr').count()} lignes)")
    B.ok("Vous : 1 astreinte" in a and pg.locator("#astBody .me").count() == 1, "Astreintes : les miennes en évidence")
    B.ok("IA 2" in a and "Rayann 2" in a, "Astreintes : répartition (majuscules confondues)")
    # suppression d'un mois
    pg.click("a[data-tab=admin]"); pg.wait_for_selector("#rcpListe [data-rcp-suppr]", state="attached")
    pg.evaluate("document.querySelector('#rcpListe [data-rcp-suppr]').click()")
    pg.wait_for_function("S.rcp.items.length === 0", timeout=8000)
    B.ok(True, "RCP d'un mois supprimées")
    # fichier qui n'est pas une fiche RCP
    pg.set_input_files("#upRcp", PLANNING); pg.wait_for_function("document.querySelector('#rcpMsg').innerText.includes('Non lu')", timeout=8000)
    B.ok("PDF" in pg.inner_text("#rcpMsg"), "fichier non PDF refusé avec un message clair")
    B.ok(not errs, "aucune erreur JavaScript" + (f" : {errs[:3]}" if errs else ""))
    b.close()
raise SystemExit(0 if B.fin() else 1)
