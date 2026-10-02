# -*- coding: utf-8 -*-
"""
Admin → Médecins : ajout, modification, désactivation et suppression d'un médecin depuis le site.
Le classeur de paramètres est modifié (onglet Médecins seulement) et reste lisible par Excel et par le générateur.
"""
import base64, io, os
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from commun import Bilan, DIST, EXEMPLE, chromium

B = Bilan("Fiche médecin dans le site")
URL = "file://" + os.path.join(DIST, "demo.html") + "#admin"
avant = load_workbook(EXEMPLE)

with sync_playwright() as p:
    b = chromium(p)
    ctx = b.new_context(viewport={"width": 390, "height": 844}, is_mobile=True, has_touch=True)
    pg = ctx.new_page(); errs = []
    pg.on("pageerror", lambda e: errs.append(str(e))); pg.on("dialog", lambda d: d.accept())
    pg.goto(URL)
    pg.wait_for_selector("#loginCard:not([hidden])")
    pg.fill("#loginEmail", "a@b.fr"); pg.fill("#loginPw", "x"); pg.click("#loginForm button[type=submit]")
    pg.wait_for_selector("#adminBody:not([hidden])")
    pg.click("#useDefault"); pg.wait_for_function("document.querySelector('#paramInfo').innerText.includes('11 médecins')")
    pg.wait_for_selector("#medList li[data-i]")
    B.ok(pg.locator("#medList li[data-i]").count() == 11, "liste des 11 médecins du classeur")

    # --- ajout
    pg.click("#medAdd"); pg.wait_for_selector("#medOverlay:not([hidden])")
    pg.fill("#mIni", "zz"); pg.fill("#mNom", "Dr Nouveau"); pg.select_option("#mQuot", "80")
    pg.click('#mOffA button[data-k="2"]'); pg.click('#mOffB button[data-k="4"]')
    pg.click('#mIndispo button[data-k="3M"]')
    pg.click('#mComp .comp-row[data-code="SCAN"] button[data-v="pr"]')
    pg.click('#mComp .comp-row[data-code="IRM2"] button[data-v="oui"]')
    pg.select_option("#mTT", "dj"); pg.click('#mTTj button[data-k="1AM"]')
    pg.click("#mSave")
    pg.wait_for_function("document.querySelector('#paramInfo').innerText.includes('12 médecins')", timeout=10000)
    B.ok(pg.is_hidden("#medOverlay") and "ZZ enregistré" in pg.inner_text("#medMsg"), "ajout : fiche enregistrée")
    d = pg.evaluate("""async () => { const P = readParams(await readXlsx(b64ToBytes(S.params.xlsx).buffer)); const z = P.docs.ZZ;
        return { comp: z.comp, quot: z.quotite, offA: [...z.offA], offB: [...z.offB], ind: [...z.indispo], tt: z.tt instanceof Set ? [...z.tt] : z.tt, nom: z.nom }; }""")
    B.ok(d["comp"].get("SCAN") == 2 and d["comp"].get("IRM2") == 1 and "MAM" not in d["comp"], "ajout : compétences (Préféré, Oui) lues par le générateur")
    B.ok(d["quot"] == 0.8 and d["offA"] == [2] and d["offB"] == [4] and d["ind"] == ["3M"] and d["tt"] == ["1AM"] and d["nom"] == "Dr Nouveau",
         "ajout : quotité, jours off A/B, indisponibilité fixe et télétravail lus par le générateur")
    # --- doublon refusé
    pg.click("#medAdd"); pg.fill("#mIni", "DA"); pg.click('#mComp .comp-row[data-code="SCAN"] button[data-v="oui"]'); pg.click("#mSave")
    B.ok("existent déjà" in pg.inner_text("#mMsg"), "initiales déjà utilisées : refusé")
    pg.click("#mCancel")
    # --- modification + désactivation
    li = pg.locator("#medList li", has_text="SB").first
    li.locator("[data-med]").click(); pg.wait_for_selector("#medOverlay:not([hidden])")
    pg.uncheck("#mActif"); pg.click("#mSave")
    pg.wait_for_function("document.querySelector('#paramInfo').innerText.includes('11 médecins')", timeout=10000)
    pg.wait_for_function("[...document.querySelectorAll('#medList li')].some((li) => li.innerText.startsWith('SB') && li.innerText.includes('inactif'))", timeout=10000)
    B.ok(True, "désactivation : SB inactif, retiré du générateur")
    # --- suppression
    pg.locator("#medList li", has_text="ZZ").first.locator("[data-med]").click(); pg.wait_for_selector("#medOverlay:not([hidden])")
    pg.click("#mDel")
    pg.wait_for_function("document.querySelector('#medMsg').innerText.includes('supprimé')", timeout=10000)
    pg.wait_for_function("![...document.querySelectorAll('#medList li')].some((li) => li.innerText.startsWith('ZZ'))", timeout=10000)
    B.ok(True, "suppression : ZZ retiré du classeur")
    # --- le classeur reste complet et lisible par Excel
    x = pg.evaluate("S.params.xlsx")
    wb = load_workbook(io.BytesIO(base64.b64decode(x)))
    B.ok(wb.sheetnames == avant.sheetnames, "classeur : tous les onglets conservés")
    B.ok([c.value for c in wb["Postes"][5]] == [c.value for c in avant["Postes"][5]], "classeur : onglet Postes inchangé")
    me = wb["Médecins"]
    B.ok(me["A1"].value == avant["Médecins"]["A1"].value and me["A4"].value == "Initiales", "classeur : titre et en-têtes conservés")
    B.ok(me["D5"].number_format == avant["Médecins"]["D5"].number_format, "classeur : mise en forme des lignes conservée")
    B.ok(len(me.data_validations.dataValidation) == len(avant["Médecins"].data_validations.dataValidation), "classeur : listes déroulantes conservées")
    sb = next(r for r in me.iter_rows(min_row=5, values_only=True) if r[0] == "SB")
    B.ok(sb[8] == "Non", "classeur : SB marqué Actif = Non")
    # --- génération possible avec le classeur modifié
    pg.select_option("#gQual", "10"); pg.click("#gBtn")
    pg.wait_for_function("document.querySelector('#gMsg').innerText.startsWith('Planning généré')", timeout=60000)
    B.ok("SB" not in pg.inner_text("#draftGrid"), "génération : médecin inactif non planifié")
    B.ok(not errs, "aucune erreur JavaScript" + (f" : {errs[:3]}" if errs else ""))
    b.close()
raise SystemExit(0 if B.fin() else 1)
