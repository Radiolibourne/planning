# -*- coding: utf-8 -*-
"""
Import du planning Excel « fait à la main » (un onglet par semaine, lignes = postes, colonnes = lundi matin … vendredi après-midi),
avec des médecins fictifs : cellules fusionnées, « * » = télétravail, Blaye à deux noms, absences « (mat) », astreinte, notes,
internes ignorés, jour férié.
"""
import datetime as dt, json, os
from openpyxl import Workbook
from playwright.sync_api import sync_playwright
from commun import Bilan, DIST, SORTIE, chromium

B = Bilan("Import du planning fait à la main")
wb = Workbook(); wb.remove(wb.active)
LIB = {5: "INTERPRETATION RADIO LIBOURNE URGENCES :", 7: "DEPLACEMENT BLAYE", 9: "DEPLACEMENT STE FOY", 10: "SCANNER : LIBOURNE, BLAYE, STE FOY",
       12: "+", 14: "INTERVENTIONNEL SCANNER", 15: "IRM 1 + RADIO URGENCES", 16: "IRM 2 + RADIO EXTERNES", 19: "SALLE  Interventionnel",
       20: "ECHOGRAPHIE 1", 21: "ECHOGRAPHIE 2", 22: "MAMMOGRAPHIE", 23: "NOTE", 26: "Absence : Rayann", 32: "ASTREINTE"}
def semaine(titre, entete, cases, fusions=()):
    ws = wb.create_sheet(titre)
    ws["A1"] = entete
    for r, t in LIB.items(): ws.cell(r, 1, t)
    for (r, c), v in cases.items(): ws.cell(r, c, v)
    for f in fusions: ws.merge_cells(f)
# colonnes : B = lun M, C = lun AM, D = mar M, … , J = ven M, K = ven AM
semaine("DU 2 NOVEMBRE AU 6 NOVEMBRE", "SEMAINE DU 2 NOVEMBRE 2026 AU 6 NOVEMBRE 2026", {
    (5, 2): "SCAN/IRM (V/S)",                       # ligne d'interprétation : ignorée
    (7, 2): "SA / OA",                              # Blaye fusionné sur la journée : 1er mammo, 2e radio-écho
    (10, 2): "DA/OB*/ INTERNE", (10, 3): "DA/SC", (10, 6): "DB/OC",
    (19, 8): "OA",                                  # salle le jeudi matin -> arthrographies
    (15, 2): "SD", (15, 3): "RAYANN",
    (23, 4): "Réunion de service à 13h30", (23, 8): "MDP Bureau", (24, 8): "Staff IRM à 13h30",
    (31, 8): "IB(mat)/DA/ SB(off à récup : le 19/11)",
    (32, 2): "IA", (32, 8): "OB", (32, 10): "Rayann",
}, fusions=("B7:C7", "D23:E23", "F10:F11", "H31:I31"))
semaine("DU 9 NOVEMBRE AU 13 NOVEMBRE", "SEMAINE DU 9 NOVEMBRE 2026 AU 13 NOVEMBRE 2026", {
    (10, 2): "DA/OB/SC", (10, 6): "FERIE", (22, 6): "FERIE",
})
f = os.path.join(SORTIE, "planning_fait_main.xlsx")
wb.save(f)
jn = lambda m, d: (dt.date(2026, m, d) - dt.date(1970, 1, 1)).days

with sync_playwright() as p:
    b = chromium(p)
    pg = b.new_context(viewport={"width": 1200, "height": 900}).new_page(); errs = []
    pg.on("pageerror", lambda e: errs.append(str(e))); pg.on("dialog", lambda d: d.accept())
    pg.goto("file://" + os.path.join(DIST, "demo.html") + "#admin")
    pg.wait_for_selector("#loginCard:not([hidden])")
    pg.fill("#loginEmail", "a@b.fr"); pg.fill("#loginPw", "x"); pg.click("#loginForm button[type=submit]")
    pg.wait_for_selector("#adminBody:not([hidden])")
    pg.click("#useDefault"); pg.wait_for_function("document.querySelector('#paramInfo').innerText.includes('médecins')")
    pg.set_input_files("#upPlanning", f)
    pg.wait_for_function("document.querySelector('#gMsg').innerText.startsWith('Planning importé')", timeout=10000)
    msg = pg.inner_text("#gMsg")
    B.ok("fait à la main" in msg and "INTERNE" in msg and "RAYANN" in msg and "11)" not in msg.split("Ignorés")[-1], f"import accepté, internes signalés comme ignorés ({msg})")
    r = pg.evaluate("""([l, j, me, ve, l2, je2]) => { const d = S.draft, c = (dd, h, k) => d.cases[caseKey(dd, h, k)] || "";
      return { start: d.start, end: d.end, ferie: d.jours.includes(me + 7),
        blm: [c(l, "M", "BL-MAM"), c(l, "AM", "BL-MAM")], blr: [c(l, "M", "BL-RE"), c(l, "AM", "BL-RE")],
        scan: c(l, "M", "SCAN"), tt: c(l, "M", "SCAN-TT"), mer: [c(me, "M", "SCAN"), c(me, "AM", "SCAN")], arth: c(j, "M", "ARTH"), si: c(j, "M", "SI"),
        irm: [c(l, "M", "IRM1"), c(l, "AM", "IRM1")], ib: [(d.statuts.IB || {})[j + "_M"], (d.statuts.IB || {})[j + "_AM"]],
        da: [(d.statuts.DA || {})[j + "_M"], (d.statuts.DA || {})[j + "_AM"]], sb: (d.statuts.SB || {})[j + "_M"],
        notes: d.notes, ordre: notesDuJour(d, j), scan2: c(l2, "M", "SCAN"), ignores: d.ignores }; }""",
        [jn(11, 2), jn(11, 5), jn(11, 4), jn(11, 6), jn(11, 9), jn(11, 12)])
    B.ok(r["start"] == jn(11, 2) and r["end"] == jn(11, 13) and not r["ferie"], "période lue dans les onglets, 11 novembre (férié) exclu")
    B.ok(r["blm"] == ["SA", "SA"] and r["blr"] == ["OA", "OA"], f"Blaye : cellule fusionnée sur la journée, 1er nom en mammo, 2e en radio-écho ({r['blm']} {r['blr']})")
    B.ok(r["scan"] == "DA" and r["tt"] == "OB", f"« * » : télétravail sur le scanner ({r['scan']} / {r['tt']})")
    B.ok(r["mer"] == ["DB / OC", ""], f"cellule fusionnée sur deux lignes : lue une seule fois ({r['mer']})")
    B.ok(r["arth"] == "OA" and r["si"] == "", "salle du jeudi matin : arthrographies")
    B.ok(r["irm"] == ["SD", ""], "interne ignoré (RAYANN)")
    B.ok(r["ib"] == ["ABS", None] and r["da"] == ["ABS", "ABS"] and r["sb"] == "ABS", f"absences : « (mat) » = le matin seulement, parenthèses ignorées ({r['ib']} {r['da']})")
    n = r["notes"] or {}
    B.ok("Réunion de service à 13h30" in n.get(str(jn(11, 3)), []) and "Astreinte : IA" in n.get(str(jn(11, 2)), []) and "Astreinte : OB" in n.get(str(jn(11, 5)), []) and "Astreinte : Rayann" in n.get(str(jn(11, 6)), [])
         and r["ordre"][:1] == ["Astreinte : OB"] and "Staff IRM à 13h30" in n.get(str(jn(11, 5)), []) and not any("Bureau" in t for t in sum(n.values(), [])),
         "notes et astreintes reprises")
    B.ok(r["scan2"] == "DA / OB / SC" and r["ignores"] is None, "2e semaine lue ; liste des ignorés non enregistrée dans le planning")
    pg.click("#publishBtn"); pg.click("#pubOk"); pg.wait_for_function("document.querySelector('#gMsg').innerText.startsWith('Planning publié')", timeout=10000)
    pg.wait_for_function("S.published && S.published.cases", timeout=10000)
    B.ok(pg.evaluate("S.published.cases[caseKey(%d, 'M', 'SCAN-TT')]" % jn(11, 2)) == "OB", "publication du planning importé")
    # retouche sur ordinateur : glisser DA (scanner) sur SD (IRM 1) = échange ; republication
    pg.click("#editPublished"); pg.wait_for_selector("#draftGrid td[data-k] .dz", timeout=8000)
    src = pg.locator(f'#draftGrid td[data-k="{jn(11, 2)}_M_SCAN"] .dz[data-ini="DA"]')
    dst = pg.locator(f'#draftGrid td[data-k="{jn(11, 2)}_M_IRM1"] .dz[data-ini="SD"]')
    B.ok(src.count() == 1 and dst.count() == 1, "ordinateur : initiales glissables dans la grille")
    src.drag_to(dst)
    g = pg.evaluate("(d) => [S.draft.cases[caseKey(d, 'M', 'SCAN')], S.draft.cases[caseKey(d, 'M', 'IRM1')], S.annuler.length]", jn(11, 2))
    B.ok(g[0] == "SD" and g[1] == "DA" and g[2] == 1, f"glisser un médecin sur un autre : échange ({g})")
    pg.click("#publishBtn"); pg.click("#pubOk"); pg.wait_for_function("S.published && S.published.publication && S.published.publication.modifs.DA", timeout=10000)
    pm = pg.evaluate("[S.published.publication, Object.keys(S.published.modifs || {}).sort()]")
    B.ok(sorted(pm[0]["modifs"]) == ["DA", "SD"] and pm[0]["nouveauxJours"] == 0 and pm[1] == ["DA", "SD"], f"republication : changements par médecin enregistrés ({pm[0]['modifs']})")
    pg.evaluate("S.me = 'DA'; S.monMois = 202611; location.hash = '#mon'"); pg.wait_for_selector("#vMon .tag.modif", timeout=8000)
    B.ok(pg.locator("#vMon .tag.modif").count() == 1, "Mon planning : case modifiée signalée")
    pg.evaluate("location.hash = '#admin'"); pg.wait_for_function("document.querySelector('#jourList').innerText.includes('demi-journée(s) modifiée(s)')", timeout=8000)
    B.ok("Planning publié" in pg.inner_text("#jourList") and "DA, SD" in pg.inner_text("#jourList"), "journal : publication et médecins concernés")
    # 2e publication (semaine suivante) : la première reste en ligne
    wb2 = Workbook(); wb2.remove(wb2.active)
    globals()["wb"] = wb2
    semaine("DU 16 NOVEMBRE AU 20 NOVEMBRE", "SEMAINE DU 16 NOVEMBRE 2026 AU 20 NOVEMBRE 2026", {(10, 2): "OA/OC/SD"})
    f2 = os.path.join(SORTIE, "planning_fait_main_2.xlsx"); wb2.save(f2)
    pg.set_input_files("#upPlanning", f2)
    pg.wait_for_function("document.querySelector('#gMsg').innerText.startsWith('Planning importé')", timeout=10000)
    pg.click("#publishBtn"); pg.click("#pubOk"); pg.wait_for_function("S.published && S.published.jours.includes(%d)" % jn(11, 16), timeout=10000)
    q = pg.evaluate("({ j: S.published.jours, tt: S.published.cases[caseKey(%d, 'M', 'SCAN-TT')], s: S.published.cases[caseKey(%d, 'M', 'SCAN')], t: S.published.titre })" % (jn(11, 2), jn(11, 16)))
    B.ok(jn(11, 2) in q["j"] and jn(11, 16) in q["j"] and q["tt"] == "OB" and q["s"] == "OA / OC / SD" and q["t"].endswith("novembre 2026"),
         f"publier une autre période conserve les semaines déjà publiées ({len(q['j'])} jours, {q['t']})")
    # Mon planning : un mois à la fois
    wb3 = Workbook(); wb3.remove(wb3.active); globals()["wb"] = wb3
    semaine("DU 7 DECEMBRE AU 11 DECEMBRE", "SEMAINE DU 7 DECEMBRE 2026 AU 11 DECEMBRE 2026", {(10, 2): "OA/OC/SD"})
    f3 = os.path.join(SORTIE, "planning_fait_main_3.xlsx"); wb3.save(f3)
    pg.set_input_files("#upPlanning", f3)
    pg.wait_for_function("document.querySelector('#gMsg').innerText.startsWith('Planning importé')", timeout=10000)
    pg.click("#publishBtn"); pg.click("#pubOk"); pg.wait_for_function("S.published && S.published.jours.includes(%d)" % jn(12, 7), timeout=10000)
    pg.evaluate("S.me = 'OA'; location.hash = '#mon'"); pg.wait_for_selector("#monMois", timeout=8000)
    opts = pg.evaluate("[...document.querySelectorAll('#monMois option')].map((o) => o.textContent)")
    pg.select_option("#monMois", label="décembre 2026"); pg.wait_for_timeout(300)
    nj = pg.locator("#vMon .day").count()
    B.ok(opts == ["novembre 2026", "décembre 2026"] and nj == 5, f"Mon planning : menu des mois, décembre seul affiché ({opts}, {nj} jours)")
    # Général, vue « Semaine » : postes en lignes, jours en colonnes, astreinte en tête
    pg.evaluate("S.genView.mode = 'semaine'; S.genView.week = String(%d); location.hash = '#general'" % jn(11, 2))
    pg.wait_for_selector("#genBody table.sem", timeout=8000)
    sem = pg.evaluate("""() => ({ ths: [...document.querySelectorAll('#genBody table.sem thead th')].map((x) => x.innerText.split('\\n')[0]),
      ast: (document.querySelector('#genBody tr.ast') || {}).innerText || '', tt: document.querySelectorAll('#genBody table.sem i.tt').length })""")
    B.ok(sem["ths"] == ["Poste", "Lun", "Mar", "Mer", "Jeu", "Ven"] and "IA" in sem["ast"] and sem["tt"] > 0, f"Général, vue Semaine ({sem})")
    # archives : un document par mois publié, conservé ; statistiques sur les archives
    pg.wait_for_function("(S.archives || []).length >= 2", timeout=8000)
    arch = pg.evaluate("S.archives.map((a) => a.id).sort()")
    B.ok(arch[:2] == ["2026-11", "2026-12"], f"archives mensuelles créées à la publication ({arch})")
    pg.evaluate("""() => { const a = extraireMois(S.published, '2026-11'); const j = (d) => d - 153;
        const vieux = { ...a, mois: '2026-06', jours: a.jours.map(j), start: j(a.start), end: j(a.end), semaines: a.semaines.map((w) => ({ ...w, lundi: j(w.lundi) })),
          cases: Object.fromEntries(Object.entries(a.cases).map(([k, v]) => { const [d, ...r] = k.split('_'); return [[j(+d), ...r].join('_'), v]; })), statuts: {}, notes: {} };
        return S.store.set(PATHS.archives() + '/2026-06', vieux); }""")
    pg.evaluate("location.hash = '#admin'"); pg.wait_for_function("[...document.querySelectorAll('#statPer option')].some((o) => o.textContent === 'juin 2026')", timeout=8000)
    B.ok(True, "statistiques : un mois ancien (archivé, plus dans le planning publié) reste disponible")
    # sauvegarde complète
    with pg.expect_download() as dl:
        pg.click("#sauvBtn")
    import zipfile, io
    chemin_z = os.path.join(SORTIE, "sauvegarde.zip"); dl.value.save_as(chemin_z)
    z = zipfile.ZipFile(chemin_z); noms = z.namelist()
    from openpyxl import load_workbook as _lw2
    feuilles = _lw2(io.BytesIO(z.read("plannings_par_mois.xlsx"))).sheetnames
    sv = json.loads(z.read("sauvegarde.json"))
    B.ok("sauvegarde.json" in noms and any(n.endswith(".xlsx") and "parametres" in n for n in noms) and feuilles[:3] == ["2026-06", "2026-11", "2026-12"] and len(sv["archives"]) >= 3,
         f"sauvegarde : données, paramètres et plannings par mois ({noms}, {feuilles})")
    pg.wait_for_function("S.sauvegarde && S.sauvegarde.derniere", timeout=8000)
    B.ok("Dernière sauvegarde le" in pg.inner_text("#sauvEtat"), "date de la dernière sauvegarde affichée")
    # statistiques et export par période
    pg.evaluate("location.hash = '#admin'"); pg.wait_for_selector("#statCard:not([hidden]) #statPer option", state="attached", timeout=8000)
    pg.select_option("#statPer", label="décembre 2026"); pg.wait_for_function("document.querySelector('#statSub').innerText.startsWith('Du 07/12/2026 au 11/12/2026')", timeout=5000)
    tout = pg.evaluate("S.published.jours.length")
    B.ok("(5 jours ouvrés)" in pg.inner_text("#statSub") and pg.input_value("#statDu") == "2026-12-07", f"statistiques : un mois choisi ({pg.inner_text('#statSub')[:60]})")
    pg.fill("#statDu", "2026-11-03"); pg.fill("#statAu", "2026-11-05"); pg.dispatch_event("#statAu", "change")
    pg.wait_for_function("document.querySelector('#statSub').innerText.includes('(3 jours ouvrés)')", timeout=5000)
    B.ok(pg.input_value("#statPer") == "perso", "statistiques : dates choisies")
    with pg.expect_download() as dl:
        pg.click("#statXlsx")
    B.ok(dl.value.suggested_filename == "statistiques_2026-11-03_au_2026-11-05.xlsx", f"export Excel de la période ({dl.value.suggested_filename})")
    from openpyxl import load_workbook as _lw
    chemin = os.path.join(SORTIE, "periode.xlsx"); dl.value.save_as(chemin)
    wbp = _lw(chemin); ws = wbp["Statistiques"]
    B.ok(wbp.sheetnames == ["Statistiques", "Couverture", "Contrôles"] and "Du 03/11/2026 au 05/11/2026" in str(ws["A2"].value) and ws["A4"].value == "Médecin",
         f"Excel des statistiques seules, pour la période ({wbp.sheetnames})")
    noms = [ws.cell(r, 1).value for r in range(5, ws.max_row + 1)]
    B.ok("DA" in noms and noms[-1] == "Total", f"une ligne par médecin et un total ({noms})")
    v = pg.evaluate("""() => { const a = { ...S.published, jours: [100, 101], start: 100, end: 101, cases: { '100_M_SCAN': 'DA' }, statuts: {}, semaines: [{ lundi: mondayOf(100), type: 'A' }] };
        const n = fusionnerPublies(a, S.published, todayDay()); return n.jours.includes(100); }""")
    B.ok(v is False, "jours publiés très anciens retirés")
    B.ok(not errs, "aucune erreur JavaScript" + (f" : {errs[:3]}" if errs else ""))
    b.close()
raise SystemExit(0 if B.fin() else 1)
