# -*- coding: utf-8 -*-
"""
Admin → Règles et postes : règles chiffrées, postes, sites, fermetures, affectations imposées, fériés
modifiés depuis le site ; le classeur reste lisible par Excel et par le générateur.
Essai sur un classeur sans onglet Fermetures (ancien format) : l'onglet est créé.
"""
import base64, datetime as dt, io, os
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from commun import Bilan, DIST, EXEMPLE, SORTIE, chromium, mois_suivant

B = Bilan("Règles et postes dans le site")
wb0 = load_workbook(EXEMPLE)
del wb0["Fermetures"]
ancien = os.path.join(SORTIE, "parametres_sans_fermetures.xlsx")
wb0.save(ancien)
d1, d2 = mois_suivant()
lundi = d1 + dt.timedelta(days=(7 - d1.weekday()) % 7)

LIRE = """async () => { const P = readParams(await readXlsx(b64ToBytes(S.params.xlsx).buffer));
  return { ttMaxDay: P.ttMaxDay, refs: P.sciRefs, refA: P.refA, ech2: [...P.postes.ECH2.opens], ech2p: P.postes.ECH2.prioMin,
    irm3: P.postes.IRM3 ? { label: P.postes.IRM3.label, site: P.postes.IRM3.site, opens: [...P.postes.IRM3.opens], remote: P.postes.IRM3.remote, need: P.postes.IRM3.need } : null,
    blaye: P.sites.Blaye.M, closures: P.closures, forced: P.forced, feries: [...P.feries.keys()] }; }"""

with sync_playwright() as p:
    b = chromium(p)
    ctx = b.new_context(viewport={"width": 390, "height": 844}, is_mobile=True, has_touch=True)
    pg = ctx.new_page(); errs = []
    pg.on("pageerror", lambda e: errs.append(str(e))); pg.on("dialog", lambda d: d.accept())
    pg.goto("file://" + os.path.join(DIST, "demo.html") + "#admin")
    pg.wait_for_selector("#loginCard:not([hidden])")
    pg.fill("#loginEmail", "a@b.fr"); pg.fill("#loginPw", "x"); pg.click("#loginForm button[type=submit]")
    pg.wait_for_selector("#adminBody:not([hidden])")
    pg.set_input_files("#upParams", ancien)
    pg.wait_for_selector("#regList [data-reg]")
    ouvrir = lambda k: (pg.click(f'#regList [data-reg="{k}"]'), pg.wait_for_selector("#regOverlay:not([hidden])"))
    attendre = lambda: pg.wait_for_function("['Enregistré', 'Règles enregistrées'].some((t) => document.querySelector('#regFormMsg').innerText.startsWith(t))", timeout=10000)

    # --- règles
    ouvrir("Règles")
    pg.fill("#rg2", (lundi + dt.timedelta(days=1)).isoformat()); pg.click("#regForm button[type=submit]")
    B.ok("doit être un lundi" in pg.inner_text("#regFormMsg"), "règles : lundi de référence vérifié")
    pg.fill("#rg2", lundi.isoformat())
    pg.fill("#rg10", "2")
    pg.click('#rg7 button[data-k="OA"]')
    pg.click("#regForm button[type=submit]"); attendre()
    pg.wait_for_timeout(300)
    r = pg.evaluate(LIRE)
    B.ok(r["ttMaxDay"] == 2 and "OA" in r["refs"] and "IA" in r["refs"], "règles : télétravail par jour et référents lus par le générateur")
    B.ok(r["refA"] == (lundi - dt.date(1970, 1, 1)).days, "règles : lundi de référence enregistré")
    pg.click("#regClose")

    # --- postes : modification et ajout
    ouvrir("Postes")
    pg.locator("#regItems li", has_text="ECH2").locator("[data-edit]").click()
    pg.click('#regForm .chips.pick button[data-k="4"]'); pg.click('#regForm .chips.pick button[data-k="5"]')   # Mer M, Mer AM
    pg.select_option("#rf6", "7")
    pg.click("#regForm button[type=submit]"); attendre()
    pg.click("#regAdd")
    pg.fill("#rf0", "irm3"); pg.fill("#rf1", "IRM 3 T"); pg.select_option("#rf2", "Libourne"); pg.fill("#rf3", "IRM")
    pg.fill("#rf4", "1"); pg.fill("#rf5", "1")
    for k in ("0", "1", "2"):
        pg.click(f'#regForm .chips.pick button[data-k="{k}"]')
    pg.select_option("#rf11", "Oui")   # télétravail
    pg.click("#regForm button[type=submit]"); attendre()
    pg.wait_for_timeout(300)
    r = pg.evaluate(LIRE)
    B.ok("2M" not in r["ech2"] and "2AM" not in r["ech2"] and "1M" in r["ech2"] and r["ech2p"] == 7, "postes : ouverture et priorité modifiées")
    B.ok(r["irm3"] and r["irm3"]["label"] == "IRM 3 T" and sorted(r["irm3"]["opens"]) == ["0AM", "0M", "1M"] and r["irm3"]["remote"], "postes : nouveau poste créé (ouverture, télétravail)")
    pg.click("#regClose")
    pg.click("#medAdd"); pg.wait_for_selector("#medOverlay:not([hidden])")
    B.ok(pg.locator('#mComp .comp-row[data-code="IRM3"]').count() == 1, "fiche médecin : le nouveau poste apparaît dans les compétences")
    pg.click("#mCancel")

    # --- sites
    ouvrir("Sites")
    pg.locator("#regItems li", has_text="Blaye").locator("[data-edit]").click()
    pg.fill("#rf2", "09:00"); pg.click("#regForm button[type=submit]"); attendre()
    pg.wait_for_timeout(300)
    B.ok(pg.evaluate(LIRE)["blaye"][0] == 540, "sites : horaire de Blaye modifié")
    pg.click("#regClose")

    # --- fermetures (onglet créé), imposées, fériés
    ouvrir("Fermetures"); pg.click("#regAdd")
    pg.select_option("#rf0", "IRM1"); pg.fill("#rf1", d1.isoformat()); pg.fill("#rf2", (d1 + dt.timedelta(days=4)).isoformat()); pg.fill("#rf4", "Remplacement")
    pg.click("#regForm button[type=submit]"); attendre(); pg.wait_for_timeout(300)
    c = pg.evaluate(LIRE)["closures"]
    B.ok(len(c) == 1 and c[0]["code"] == "IRM1" and c[0]["d2"] - c[0]["d1"] == 4, "fermetures : onglet créé et fermeture lue par le générateur")
    pg.click("#regClose")
    ouvrir("Imposées"); pg.click("#regAdd")
    pg.fill("#rf0", d1.isoformat()); pg.select_option("#rf1", "Après-midi"); pg.select_option("#rf2", "OA"); pg.select_option("#rf3", "IRM2")
    pg.click("#regForm button[type=submit]"); attendre(); pg.wait_for_timeout(300)
    f = pg.evaluate(LIRE)["forced"]
    B.ok(any(x["ini"] == "OA" and x["p"] == "IRM2" and x["h"] == "AM" for x in f), "imposées : affectation lue par le générateur")
    pg.click("#regClose")
    ouvrir("Fériés")
    n0 = len(pg.evaluate(LIRE)["feries"])
    pg.locator("#regItems li", has_text="Toussaint").locator("[data-edit]").click(); pg.click("#regDel")
    pg.wait_for_selector("#regAdd"); pg.click("#regAdd")
    pg.fill("#rf0", "2027-12-26"); pg.fill("#rf1", "Lendemain de Noël")
    pg.click("#regForm button[type=submit]"); attendre(); pg.wait_for_timeout(300)
    fe = pg.evaluate(LIRE)["feries"]
    B.ok(len(fe) == n0 and (dt.date(2027, 12, 26) - dt.date(1970, 1, 1)).days in fe, "fériés : suppression et ajout")
    pg.click("#regClose")

    # --- classeur lisible par Excel, génération possible
    wb = load_workbook(io.BytesIO(base64.b64decode(pg.evaluate("S.params.xlsx"))))
    B.ok("Fermetures" in wb.sheetnames and wb["Fermetures"]["A4"].value == "Code poste", "classeur : onglet Fermetures ajouté, lisible par Excel")
    B.ok(wb["Postes"]["A4"].value == "Code" and wb["Paramètres"]["A4"].value == wb0["Paramètres"]["A4"].value, "classeur : en-têtes conservés")
    pg.fill("#gStart", d1.isoformat()); pg.fill("#gEnd", d2.isoformat()); pg.select_option("#gQual", "10"); pg.click("#gBtn")
    pg.wait_for_function("document.querySelector('#gMsg').innerText.startsWith('Planning généré')", timeout=60000)
    B.ok(True, "génération avec le classeur modifié")
    B.ok(not errs, "aucune erreur JavaScript" + (f" : {errs[:3]}" if errs else ""))
    b.close()
raise SystemExit(0 if B.fin() else 1)
