# -*- coding: utf-8 -*-
"""
Scénario complet en mode démonstration : un médecin et l'administrateur dans deux onglets.
Produit aussi (via l'assistant de mise en ligne) la page configurée et les règles utilisées par test_firebase.py.
"""
import datetime as dt
import os
import time
import urllib.parse
from playwright.sync_api import sync_playwright
from commun import Bilan, DIST, SORTIE, chromium, jours_ouvres, mois_suivant, numero_jour, verifier_dist

MEDECIN = "SB"          # médecin fictif (off le mercredi)


def main():
    verifier_dist()
    B = Bilan("Mode démonstration")
    url = "file://" + os.path.join(DIST, "demo.html")
    d1, d2 = mois_suivant()
    ouvres = jours_ouvres(d1, d2)
    # congés : la 2e semaine complète du mois (lundi → vendredi)
    lundis = [d for d in ouvres if d.weekday() == 0]
    conge_debut = lundis[1] if len(lundis) > 1 else lundis[0]
    conge_fin = conge_debut + dt.timedelta(days=4)
    erreurs_js, reseau = [], []
    with sync_playwright() as p:
        b = chromium(p)
        ctx = b.new_context(accept_downloads=True, viewport={"width": 390, "height": 844}, is_mobile=True, has_touch=True)
        ctx.route("**/*", lambda r: r.continue_() if r.request.url.startswith(("file:", "blob:", "data:"))
                  else (reseau.append(r.request.url), r.abort()))
        doc, adm = ctx.new_page(), ctx.new_page()
        for pg in (doc, adm):
            pg.on("pageerror", lambda e: erreurs_js.append(str(e)))
            pg.on("dialog", lambda d: d.accept())
        # --- médecin, avant tout paramétrage
        doc.goto(url)
        doc.wait_for_selector("#vMon:not([hidden])")
        B.ok(doc.is_visible("#demoBanner"), "bandeau démonstration visible")
        doc.click("#meBtn")
        doc.wait_for_selector("#whoOverlay:not([hidden])")
        B.ok("pas encore chargé" in doc.inner_text("#whoGrid"), "aucun médecin tant que les paramètres ne sont pas chargés")
        doc.click("#whoClose")
        # --- administrateur : paramètres
        adm.goto(url + "#admin")
        adm.wait_for_selector("#loginCard:not([hidden])")
        adm.fill("#loginEmail", "admin@test.fr"); adm.fill("#loginPw", "x"); adm.click("#loginForm button[type=submit]")
        adm.wait_for_selector("#adminBody:not([hidden])")
        adm.click("#useDefault")
        adm.wait_for_function("document.querySelector('#paramInfo').innerText.includes('médecins')")
        B.ok("11 médecins" in adm.inner_text("#paramInfo"), "paramètres d'exemple chargés")
        # --- médecin : initiales + congés
        doc.click("#meBtn"); doc.wait_for_selector(f"#whoGrid button[data-ini={MEDECIN}]", timeout=5000)
        doc.click(f"#whoGrid button[data-ini={MEDECIN}]")
        B.ok(doc.inner_text("#meBtn") == f"Vous : {MEDECIN}", "initiales mémorisées")
        doc.click("a[data-tab=indispos]")
        doc.fill("#indD1", conge_debut.isoformat()); doc.fill("#indD2", conge_fin.isoformat())
        doc.select_option("#indMotif", "Congés"); doc.fill("#indPrec", "vacances")
        doc.click("#indSubmit")
        doc.wait_for_function("document.querySelector('#indMsg').innerText.includes('en attente de validation')")
        doc.wait_for_function(f"document.querySelector('#myInd').innerText.includes('{conge_debut:%d/%m}')")
        B.ok("Congés — vacances" in doc.inner_text("#myInd") and "en attente" in doc.inner_text("#myInd"), "demande d'absence listée « en attente »")
        doc.fill("#indD1", conge_fin.isoformat()); doc.fill("#indD2", conge_debut.isoformat()); doc.click("#indSubmit")
        B.ok("avant la date" in doc.inner_text("#indMsg"), "dates incohérentes refusées")
        # --- administrateur : génération sur le mois suivant
        adm.wait_for_function(f"document.querySelector('#admInd').innerText.includes('{MEDECIN}')")
        B.ok(True, "l'administrateur voit la demande en temps réel")
        B.ok("Admin · 1" in adm.inner_text("#adminBtn"), "administrateur : nombre de demandes à valider sur le bouton Admin")
        n_abs = adm.evaluate("() => { const P = { ...S.P, abs: S.P.abs.concat(onlineAbs(S.indispos)) }; return onlineAbs(S.indispos).length; }")
        B.ok(n_abs == 0, "demande en attente : pas prise en compte par la génération")
        adm.click("#admInd [data-accept]")
        doc.wait_for_function("document.querySelector('#myInd').innerText.includes('acceptée')", timeout=5000)
        adm.wait_for_function("document.querySelector('#adminBtn').innerText === 'Admin ✓'", timeout=5000)
        B.ok(True, "administrateur : demande acceptée, le médecin le voit")
        adm.fill("#gStart", d1.isoformat()); adm.fill("#gEnd", d2.isoformat())
        adm.select_option("#gQual", "10")
        t = time.time(); adm.click("#gBtn")
        adm.wait_for_function("document.querySelector('#gMsg').innerText.startsWith('Planning généré')", timeout=90000)
        B.ok(True, f"génération en {time.time() - t:.0f} s")
        B.ok("Erreur" not in adm.inner_text("#dChecks"), "aucune erreur de règle")
        n = adm.evaluate("""([ini, a, b]) => { const d = JSON.parse(localStorage.getItem('planning-radio-draft')).draft;
            return Object.entries(d.cases).filter(([k, v]) => { const j = +k.split('_')[0];
              return j >= a && j <= b && v.split('/').map((s) => s.trim()).includes(ini); }).length; }""",
                         [MEDECIN, numero_jour(conge_debut), numero_jour(conge_fin)])
        B.ok(n == 0, f"{MEDECIN} non affecté pendant ses congés")
        # retouche manuelle
        cell = adm.locator("#draftGrid td[data-k]").first
        k, avant = cell.get_attribute("data-k"), cell.inner_text()
        cell.click(); adm.keyboard.press("Control+A"); adm.keyboard.type("da oa, zz"); adm.keyboard.press("Enter")
        B.ok(adm.locator(f'#draftGrid td[data-k="{k}"]').inner_text() == "DA / OA / ZZ", "retouche saisie et normalisée")
        B.ok("ZZ ne figure pas" in adm.inner_text("#dChecks"), "initiales inconnues signalées")
        adm.locator(f'#draftGrid td[data-k="{k}"]').click(); adm.keyboard.press("Control+A")
        adm.keyboard.type(avant if avant != "FERMÉ" else ""); adm.keyboard.press("Enter")
        # publication
        adm.click("#publishBtn")
        adm.wait_for_function("document.querySelector('#gMsg').innerText.startsWith('Planning publié')")
        B.ok(adm.is_hidden("#draftBody"), "publication effectuée")
        # --- médecin : son planning
        doc.click("a[data-tab=mon]"); doc.wait_for_selector("#vMon .day", timeout=5000)
        B.ok(doc.query_selector("#vMon #icsMe") is None, "Mon planning : plus de bouton calendrier (onglet Calendrier)")
        B.ok(doc.locator("#vMon .day").count() == len(ouvres), f"planning personnel : {len(ouvres)} jours ouvrés")
        B.ok("Absent" in doc.inner_text("#vMon"), "congés affichés")
        # semaine en cours en premier, semaines passées repliées (aujourd'hui simulé au mercredi de la 2e semaine)
        w2 = doc.evaluate("S.published.semaines[1].lundi")
        jour = dt.date(1970, 1, 1) + dt.timedelta(days=w2 + 2)
        p2 = doc.context.new_page()   # horloge simulée pour cette page seulement
        ms = int(dt.datetime(jour.year, jour.month, jour.day, 10, 0).timestamp() * 1000)
        p2.add_init_script("(() => { const T = %d, D = Date; class F extends D { constructor(...a) { a.length ? super(...a) : super(T); } static now() { return T; } } window.Date = F; })();" % ms)
        p2.goto(doc.url.split("#")[0] + "#mon"); p2.wait_for_selector("#vMon .week-h", timeout=8000)
        r = {"premiere": p2.inner_text("#vMon .week-h"), "passees": p2.locator("#vMon details.passees .week-h").count()}
        p2.close()
        B.ok("cette semaine" in r["premiere"] and r["passees"] == 1, f"Mon planning : semaine en cours en premier, semaines passées repliées ({r})")
        doc.click("a[data-tab=cal]"); doc.wait_for_selector("#icsMe", timeout=5000)
        ics = urllib.parse.unquote(doc.get_attribute("#icsMe", "href").split(",", 1)[1])
        B.ok(ics.startswith("BEGIN:VCALENDAR") and ics.count("BEGIN:VEVENT") > 5, "calendrier iPhone généré")
        # verrou : période publiée fermée aux médecins
        jour = doc.evaluate("""() => { const t = document.querySelector('#vMon .day .slot:not(.off)');
            return t ? t.closest('.day').querySelector('.date span').innerText : null; }""")
        dd, mm = jour.split("/")
        an = d1.year if int(mm) >= d1.month else d2.year
        jour_travaille = f"{an}-{mm}-{dd}"
        doc.click("a[data-tab=indispos]")
        B.ok("Planning publié jusqu'au" in doc.inner_text("#indLock"), "bandeau : période publiée verrouillée")
        doc.evaluate("(v) => { const i = document.querySelector('#indD1'); i.removeAttribute('min'); i.value = v; document.querySelector('#indD2').value = v; }", jour_travaille)
        doc.click("#indSubmit")
        B.ok("publié jusqu'au" in doc.inner_text("#indMsg"), "médecin : absence refusée sur la période publiée")
        B.ok(doc.input_value("#indD1") == jour_travaille, "la date saisie n'est jamais modifiée par la page")
        # l'administrateur peut toujours en ajouter une (ex. arrêt maladie) → conflit signalé
        adm.click("#meBtn"); adm.wait_for_selector(f"#whoGrid button[data-ini={MEDECIN}]"); adm.click(f"#whoGrid button[data-ini={MEDECIN}]")
        adm.click("a[data-tab=indispos]")
        adm.fill("#indD1", jour_travaille); adm.fill("#indD2", jour_travaille)
        adm.select_option("#indMotif", "Autre"); adm.fill("#indPrec", "arrêt")
        adm.click("#indSubmit"); adm.wait_for_function("document.querySelector('#indMsg').innerText.includes('enregistrée')")
        B.ok(True, "administrateur : absence tardive enregistrée")
        doc.click("a[data-tab=mon]"); doc.wait_for_selector(".banner.warn", timeout=5000)
        B.ok("absence signalée" in doc.inner_text("#vMon"), "conflit signalé au médecin")
        doc.click("#confFermer"); doc.wait_for_timeout(300)
        B.ok(doc.locator("#confBanner").count() == 0 and "absence signalée" not in doc.inner_text("#vMon"), "bandeau de conflit refermé d'une croix")
        adm.click("a[data-tab=admin]")
        adm.wait_for_function("document.querySelector('#admInd').innerText.includes('affecté pendant')", timeout=5000)
        B.ok(True, "conflit signalé à l'administrateur")
        # dates limites de dépôt : le mois suivant la période publiée, clos depuis hier
        m2_debut = d2 + dt.timedelta(days=1)
        m2_fin = (m2_debut.replace(day=28) + dt.timedelta(days=4)).replace(day=1) - dt.timedelta(days=1)
        hier = dt.date.today() - dt.timedelta(days=1)
        adm.fill("#cDu", m2_debut.isoformat()); adm.fill("#cAu", m2_fin.isoformat()); adm.fill("#cLim", hier.isoformat())
        adm.click("#cAdd"); adm.wait_for_function("document.querySelector('#cloList').innerText.includes('clos')")
        B.ok(True, "administrateur : date limite de dépôt enregistrée")
        doc.click("a[data-tab=indispos]")
        doc.wait_for_function("document.querySelector('#indLock').innerText.includes('dépôt clos')", timeout=5000)
        B.ok(True, "bandeau : dépôt clos affiché au médecin")
        x = m2_debut + dt.timedelta(days=7)
        doc.fill("#indD1", x.isoformat()); doc.fill("#indD2", x.isoformat()); doc.click("#indSubmit")
        B.ok("est clos depuis" in doc.inner_text("#indMsg"), "médecin : absence refusée après la date limite")
        y = m2_fin + dt.timedelta(days=10)
        doc.fill("#indD1", y.isoformat()); doc.fill("#indD2", y.isoformat()); doc.click("#indSubmit")
        doc.wait_for_function("document.querySelector('#indMsg').innerText.includes('en attente')", timeout=5000)
        B.ok(True, "médecin : demande possible hors période close")
        adm.wait_for_selector("#admInd [data-refuse]", timeout=5000); adm.click("#admInd [data-refuse]")
        doc.wait_for_function("document.querySelector('#myInd').innerText.includes('refusée')", timeout=5000)
        B.ok(True, "administrateur : demande refusée, le médecin le voit")
        # périodes ouvertes : une 2e période plus loin crée un « trou » où les demandes sont fermées
        adm.click("a[data-tab=admin]")
        p3, p3_fin, demain = m2_fin + dt.timedelta(days=40), m2_fin + dt.timedelta(days=60), dt.date.today() + dt.timedelta(days=1)
        adm.fill("#cDu", p3.isoformat()); adm.fill("#cAu", p3_fin.isoformat()); adm.fill("#cLim", demain.isoformat())
        adm.click("#cAdd"); adm.wait_for_function("document.querySelectorAll('#cloList > li').length >= 2", timeout=5000)
        doc.click("a[data-tab=mon]"); doc.click("a[data-tab=indispos]")
        doc.fill("#indD1", y.isoformat()); doc.fill("#indD2", y.isoformat()); doc.click("#indSubmit")
        B.ok("ne sont pas ouvertes" in doc.inner_text("#indMsg"), "médecin : demande refusée entre deux périodes ouvertes")
        z = p3 + dt.timedelta(days=3)
        doc.fill("#indD1", z.isoformat()); doc.fill("#indD2", z.isoformat()); doc.click("#indSubmit")
        doc.wait_for_function("document.querySelector('#indMsg').innerText.includes('en attente')", timeout=5000)
        w = p3_fin + dt.timedelta(days=5)
        doc.fill("#indD1", w.isoformat()); doc.fill("#indD2", w.isoformat()); doc.click("#indSubmit")
        doc.wait_for_function("document.querySelector('#indMsg').innerText.includes('en attente')", timeout=5000)
        B.ok("Demandes possibles uniquement" in doc.inner_text("#indLock") and "à partir du" in doc.inner_text("#indLock"), "médecin : demandes possibles dans une période ouverte et après la dernière")
        # planning général
        doc.click("a[data-tab=general]"); doc.wait_for_selector("#genBody table")
        B.ok(doc.locator("#genBody td.me, #genBody .dj.me").count() > 0, "cases du médecin encadrées")
        B.ok(doc.evaluate("document.documentElement.scrollWidth <= window.innerWidth"), "pas de défilement horizontal sur mobile")
        with doc.expect_download() as d:
            doc.click("#genXlsx")
        d.value.save_as(os.path.join(SORTIE, "general.xlsx"))
        from openpyxl import load_workbook as _lw
        _f = _lw(os.path.join(SORTIE, "general.xlsx"))
        B.ok("Planning par poste" in _f.sheetnames, "médecin : planning publié téléchargé en Excel")
        B.ok(not ({"Synthèse", "Contrôles", "Lisez-moi", "Fermetures"} & set(_f.sheetnames)), f"médecin : Excel sans statistiques ni contrôles des règles {_f.sheetnames}")
        # administrateur : statistiques et Excel complet
        adm.click("a[data-tab=admin]"); adm.wait_for_selector("#statCard:not([hidden]) #statTable table", timeout=5000)
        B.ok(adm.locator("#statTable tbody tr").count() >= 8 and "Demi-j." in adm.inner_text("#statTable"), "administrateur : statistiques par médecin")
        B.ok(adm.locator("#statChecks li").count() > 0 and adm.locator("#statCov .meter").count() > 0, "administrateur : couverture et contrôle des règles du planning publié")
        with adm.expect_download() as d:
            adm.click("#statXlsx")
        d.value.save_as(os.path.join(SORTIE, "complet.xlsx"))
        B.ok(_lw(os.path.join(SORTIE, "complet.xlsx")).sheetnames == ["Statistiques", "Couverture", "Contrôles"], "administrateur : Excel des statistiques (sans les plannings)")
        # exports depuis la copie du planning publié
        adm.click("#editPublished"); adm.wait_for_selector("#draftBody:not([hidden])")
        with adm.expect_download() as d:
            adm.click("#dlDraftXlsx")
        d.value.save_as(os.path.join(SORTIE, "export.xlsx"))
        with adm.expect_download() as d:
            adm.click("#dlDraftIcs")
        d.value.save_as(os.path.join(SORTIE, "calendriers.zip"))
        # assistant de mise en ligne → page configurée + règles (réutilisées par test_firebase.py)
        adm.fill("#sCfg", 'const firebaseConfig = { apiKey: "AIzaTEST", authDomain: "test.firebaseapp.com", '
                          'projectId: "test", storageBucket: "test.appspot.com", messagingSenderId: "1", appId: "1:1:web:1" };')
        adm.click("#sGenCode"); code = adm.input_value("#sCode")
        adm.fill("#sUids", "AdminUidDeTest0000000000")
        adm.click("#sBuild"); adm.wait_for_selector("#sOut:not([hidden])")
        regles = adm.input_value("#sRules")
        B.ok(code in regles and "AdminUidDeTest0000000000" in regles, "règles de sécurité générées")
        with adm.expect_download() as d:
            adm.click("#sDl")
        d.value.save_as(os.path.join(SORTIE, "index_configure.html"))
        with open(os.path.join(SORTIE, "regles.txt"), "w", encoding="utf-8") as f:
            f.write(regles)
        b.close()
    html = open(os.path.join(SORTIE, "index_configure.html"), encoding="utf-8").read()
    B.ok('"apiKey": "AIzaTEST"' in html, "page configurée : configuration Firebase présente")
    B.ok('DEFAULT_PARAMS_B64 = ""' in html and code not in html, "page configurée : ni données du service, ni code d'équipe")
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(SORTIE, "export.xlsx"))
    B.ok({"Planning par poste", "Planning par médecin", "Synthèse"} <= set(wb.sheetnames), "export Excel lisible")
    B.ok(not reseau, "aucune connexion réseau en démonstration")
    B.ok(not erreurs_js, "aucune erreur JavaScript" + (f" : {erreurs_js[:3]}" if erreurs_js else ""))
    return B.fin()


if __name__ == "__main__":
    raise SystemExit(0 if main() else 1)
