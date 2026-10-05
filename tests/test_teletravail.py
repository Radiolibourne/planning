# -*- coding: utf-8 -*-
"""
Télétravail, fermetures ponctuelles et import d'un planning Excel (moteur testé dans la page de démonstration).
- 1 jour de télétravail par semaine au plus, si possible pour chaque médecin autorisé ;
- un jour de télétravail : seulement des postes faisables à distance, jamais mélangé avec une présence sur site ;
- médecin « Non » jamais en télétravail, médecin « Jeu AM » seulement le jeudi après-midi ;
- maximum de médecins en télétravail par jour, répartition équitable sur le mois ;
- le télétravail compte pour la couverture du poste d'origine ;
- onglet Fermetures : poste fermé sur la période ;
- classeur sans colonne Télétravail : aucun télétravail (compatibilité) ;
- un planning exporté en Excel puis réimporté redonne le même planning.
"""
import base64, datetime as dt, io, json, os
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
from commun import Bilan, DIST, EXEMPLE, SORTIE, chromium, mois_suivant, numero_jour

B = Bilan("Télétravail, fermetures, import Excel")
d1, d2 = mois_suivant()
src = load_workbook(EXEMPLE)

# 1) classeur d'exemple + fermeture de l'IRM 1 sur la 2e semaine
wb = load_workbook(EXEMPLE)
ws = wb["Paramètres"]
for r in range(5, 40):
    if ws.cell(r, 1).value == "Début de période": ws.cell(r, 2).value = d1
    if ws.cell(r, 1).value == "Fin de période": ws.cell(r, 2).value = d2
lundis = [d1 + dt.timedelta(days=i) for i in range((d2 - d1).days + 1) if (d1 + dt.timedelta(days=i)).weekday() == 0]
f1, f2 = lundis[1], lundis[1] + dt.timedelta(days=4)
wf = wb["Fermetures"]
wf.cell(5, 1, "IRM1"); wf.cell(5, 2, f1); wf.cell(5, 3, f2); wf.cell(5, 4, "Journée"); wf.cell(5, 5, "Remplacement")
buf = io.BytesIO(); wb.save(buf); avec_tt = base64.b64encode(buf.getvalue()).decode()

# 2) même classeur sans aucune colonne Télétravail (ancien format)
wb2 = load_workbook(EXEMPLE)
for r in range(5, 40):
    if wb2["Paramètres"].cell(r, 1).value == "Début de période": wb2["Paramètres"].cell(r, 2).value = d1
    if wb2["Paramètres"].cell(r, 1).value == "Fin de période": wb2["Paramètres"].cell(r, 2).value = d2
for name in ("Postes", "Médecins"):
    w = wb2[name]
    for c in range(1, w.max_column + 1):
        if str(w.cell(4, c).value or "").startswith("Télétravail"): w.delete_cols(c)
del wb2["Fermetures"]
buf = io.BytesIO(); wb2.save(buf); sans_tt = base64.b64encode(buf.getvalue()).decode()

JS = r"""
async ([b64, secs]) => {
  const bytes = Uint8Array.from(atob(b64), (c) => c.charCodeAt(0));
  const P = readParams(await readXlsx(bytes.buffer));
  const pr = buildProblem(P);
  const R = await solve(pr, secs);
  const items = checks(pr, R);
  const cov = coverage(pr, R);
  const docs = pr.docs;
  const ttOf = {}, mix = [], parSlot = {}, horsPoste = [];
  pr.slots.forEach((sl, s) => pr.postes.forEach((code, p) => {
    for (const ini of R.assign.get(s + "|" + p) || []) {
      if (!pr.pTT[p]) continue;
      (ttOf[ini] = ttOf[ini] || []).push([sl.d, sl.h, weekday(sl.d)]);
      (parSlot[sl.d] = parSlot[sl.d] || new Set()).add(ini);
      if (!P.postes[code].base || !P.postes[P.postes[code].base].remote) horsPoste.push(code);
    }
  }));
  // mélange télétravail / site le même jour
  docs.forEach((ini, d) => pr.days.forEach((day, di) => {
    const ps = pr.slotsOfDay[di].map((s) => R.a[d][s]).filter((p) => p >= 0);
    if (ps.some((p) => pr.pTT[p]) && ps.some((p) => !pr.pTT[p])) mix.push(ini + " " + day);
  }));
  // jours de télétravail par médecin et par semaine
  const parSemaine = {};
  for (const [ini, l] of Object.entries(ttOf)) {
    const m = {};
    for (const [d] of l) { const w = mondayOf(d); (m[w] = m[w] || new Set()).add(d); }
    parSemaine[ini] = Object.values(m).map((x) => x.size);
  }
  const irm1 = pr.pIdx["IRM1"];
  const irm1Ouvert = pr.slots.filter((sl, s) => pr.open[s][irm1]).map((sl) => sl.d);
  const draft = draftFromResult(pr, R);
  return { tt: P.tt, twins: pr.postes.filter((c, p) => pr.pTT[p]), ttOf, mix, parSlot: Object.values(parSlot).map((x) => x.size), horsPoste,
           parSemaine, erreurs: items.filter((i) => i[0] === "Erreur").map((i) => i[1]), cov,
           irm1Ouvert, sites: draft.sites, nWeeks: pr.weeks.length, docs, ttDocs: docs.filter((x) => P.docs[x].tt) };
}
"""

with sync_playwright() as p:
    b = chromium(p)
    ctx = b.new_context(accept_downloads=True)
    pg = ctx.new_page()
    errs = []
    pg.on("pageerror", lambda e: errs.append(str(e)))
    pg.goto("file://" + os.path.join(DIST, "demo.html"))
    pg.wait_for_function("typeof solve === 'function'")
    r = pg.evaluate(JS, [avec_tt, 8])
    B.ok(r["tt"] and sorted(r["twins"]) == ["IRM2-TT", "SCAN-TT"], f"postes de télétravail créés : {r['twins']}")
    B.ok(not r["erreurs"], "aucune erreur de règle" + (f" : {r['erreurs'][:3]}" if r["erreurs"] else ""))
    B.ok(not r["mix"], "jamais télétravail et site le même jour" + (f" : {r['mix'][:3]}" if r["mix"] else ""))
    B.ok(all(n <= 1 for v in r["parSemaine"].values() for n in v), "au plus 1 jour de télétravail par semaine")
    B.ok(all(n <= 1 for n in r["parSlot"]), "au plus 1 médecin en télétravail par jour (paramètre)")
    libres = [sum(r["parSemaine"].get(i, [])) for i in r["ttDocs"] if i not in ("IA", "IB")]
    B.ok(max(libres) - min(libres) <= 2, f"télétravail réparti équitablement sur le mois {libres}")
    B.ok("IA" not in r["ttOf"], "IA (Télétravail = Non) jamais en télétravail")
    B.ok(all(wd == 3 and h == "AM" for _, h, wd in r["ttOf"].get("IB", [])), "IB seulement le jeudi après-midi")
    B.ok(not r["horsPoste"], "télétravail seulement sur les postes faisables à distance")
    total = sum(len(v) for v in r["parSemaine"].values())
    B.ok(total >= r["nWeeks"] * 4, f"le télétravail est bien attribué ({total} jours sur {r['nWeeks']} semaines)")
    scan = next(c for c in r["cov"] if c["code"] == "SCAN")
    B.ok(not any(c["code"].endswith("-TT") for c in r["cov"]) and scan["tt"] > 0, "couverture : le télétravail compte pour le scanner")
    B.ok(all(not (numero_jour(f1) <= d <= numero_jour(f2)) for d in r["irm1Ouvert"]) and r["irm1Ouvert"], "fermeture : IRM 1 fermée sur la période indiquée")
    B.ok(r["sites"][-1] == "Télétravail", "planning publié : site « Télétravail » ajouté")
    r2 = pg.evaluate(JS, [sans_tt, 3])
    B.ok(not r2["tt"] and not r2["twins"] and not r2["ttOf"], "classeur sans colonne Télétravail : aucun télétravail")
    # export Excel puis réimport -> même planning
    rt = pg.evaluate("""async (b64) => {
      const bytes = Uint8Array.from(atob(b64), (c) => c.charCodeAt(0));
      const P = readParams(await readXlsx(bytes.buffer));
      const pr = buildProblem(P);
      const R = await solve(pr, 3);
      const d0 = draftFromResult(pr, R);
      const items = checks(pr, R);
      const out = await exportPlanning(pr, R, items);
      const blob = out instanceof Blob ? out : new Blob([out]);
      const d1 = draftFromPlanningXlsx(await readXlsx(await blob.arrayBuffer()), P);
      const diff = Object.keys(d0.cases).filter((k) => (d0.cases[k] || "") !== (d1.cases[k] || ""));
      const st = Object.keys(d0.statuts).filter((ini) => JSON.stringify(d0.statuts[ini]) !== JSON.stringify(d1.statuts[ini] || {}));
      return { diff: diff.slice(0, 5), n: diff.length, st, start: JSON.stringify(d0.jours) === JSON.stringify(d1.jours),
               sites: JSON.stringify(d0.sites) === JSON.stringify(d1.sites), ordre: JSON.stringify(d0.ordre) === JSON.stringify(d1.ordre),
               labels: d0.ordre.every((c) => d1.postes[c] && d1.postes[c].label === d0.postes[c].label && d1.postes[c].site === d0.postes[c].site) };
    }""", avec_tt)
    B.ok(rt["n"] == 0, f"import Excel : mêmes affectations ({rt['n']} différences {rt['diff']})")
    B.ok(not rt["st"], f"import Excel : mêmes repos et absences {rt['st'][:3]}")
    B.ok(rt["start"] and rt["sites"] and rt["ordre"] and rt["labels"], "import Excel : période, sites et postes identiques")
    fold = pg.evaluate("""async ([a, b]) => {
      const load = async (x) => readParams(await readXlsx(Uint8Array.from(atob(x), (c) => c.charCodeAt(0)).buffer));
      const P1 = await load(a), P2 = await load(b);
      const pr1 = buildProblem(P1), R = await solve(pr1, 3), draft = draftFromResult(pr1, R);
      const pr2 = buildProblem(P2), A = assignFromDraft(pr2, draft);
      const scan2 = pr2.pIdx["SCAN"], ttKeys = Object.keys(draft.cases).filter((k) => k.endsWith("_SCAN-TT") && draft.cases[k]);
      const ok = ttKeys.every((k) => { const [d, h] = k.split("_"); const s = pr2.slots.findIndex((x) => x.d === +d && x.h === h);
        return splitInis(draft.cases[k]).every((i) => (A.get(s + "|" + scan2) || []).includes(i)); });
      return { n: ttKeys.length, ok };
    }""", [avec_tt, sans_tt])
    B.ok(fold["n"] > 0 and fold["ok"], "planning avec télétravail contrôlé avec des paramètres sans télétravail : compté au scanner")
    # parcours réel : Admin -> Importer un planning Excel -> Publier ; repos/absences et médecins du fichier conservés
    x = pg.evaluate("""async (b64) => {
      const P = readParams(await readXlsx(Uint8Array.from(atob(b64), (c) => c.charCodeAt(0)).buffer));
      const pr = buildProblem(P), R = await solve(pr, 2);
      const k = [...R.assign.keys()][0]; R.assign.get(k).push("ZZ");          // médecin absent des paramètres
      let bin = ""; const u = new Uint8Array(await exportPlanning(pr, R, checks(pr, R)));
      for (let i = 0; i < u.length; i++) bin += String.fromCharCode(u[i]);
      return btoa(bin);
    }""", avec_tt)
    fx = os.path.join(SORTIE, "planning_a_importer.xlsx")
    open(fx, "wb").write(base64.b64decode(x))
    wbx = load_workbook(fx); stx = wbx["Statuts"]
    col = next(c for c in range(4, stx.max_column + 1) if stx.cell(5, c).value == "SA")
    lig = next(r for r in range(7, stx.max_row + 1) if stx.cell(r, col).value == "DISPO")
    stx.cell(lig, col).value = "ABS"; wbx.save(fx)                         # absence propre au fichier
    pg.on("dialog", lambda d: d.accept())
    pg.goto("file://" + os.path.join(DIST, "demo.html") + "#admin")
    pg.wait_for_selector("#loginCard:not([hidden])")
    pg.fill("#loginEmail", "a@b.fr"); pg.fill("#loginPw", "x"); pg.click("#loginForm button[type=submit]")
    pg.wait_for_selector("#adminBody:not([hidden])")
    pg.click("#useDefault"); pg.wait_for_function("document.querySelector('#paramInfo').innerText.includes('médecins')")
    pg.set_input_files("#upPlanning", fx)
    pg.wait_for_function("document.querySelector('#gMsg').innerText.startsWith('Planning importé')", timeout=10000)
    B.ok(pg.is_visible("#draftBody"), "Admin : planning Excel importé comme brouillon")
    pg.click("#publishBtn"); pg.wait_for_function("document.querySelector('#gMsg').innerText.startsWith('Planning publié')")
    pg.wait_for_function("S.published && (S.published.medecins || []).includes('ZZ')", timeout=15000)   # la mise à jour de S.published peut suivre le message
    pub = pg.evaluate("S.published")
    B.ok("ZZ" in pub["medecins"], "publication : médecin du fichier absent des paramètres conservé")
    B.ok("ABS" in pub["statuts"].get("SA", {}).values(), "publication : absence propre au fichier conservée")
    # télétravail affiché sur le poste, avec une étoile (pas de colonne « Télétravail »)
    pp = wbx["Planning par poste"]
    entetes = [str(pp.cell(4, c).value or "") for c in range(1, pp.max_column + 1)]
    etoiles = [c.value for row in pp.iter_rows(min_row=7) for c in row if isinstance(c.value, str) and c.value.endswith("*")]
    B.ok(not any("TÉLÉTRAVAIL" in e.upper() for e in entetes) and etoiles, f"Excel : télétravail sur le poste avec une étoile ({etoiles[:3]})")
    B.ok(any(k.endswith("_SCAN-TT") and v for k, v in pub["cases"].items()), "import Excel : « CD* » relu comme télétravail")
    pg.click("a[data-tab=general]"); pg.wait_for_selector("#vGeneral table.pl", timeout=8000)
    t = pg.evaluate("""() => ({ ths: [...document.querySelectorAll('#vGeneral th')].map((x) => x.innerText), tt: document.querySelectorAll('#vGeneral i.tt').length })""")
    B.ok(not any(x.strip().upper() == "TÉLÉTRAVAIL" for x in t["ths"]) and t["tt"] > 0, f"Général : médecins en télétravail sur leur poste, en italique avec étoile ({t['tt']})")
    pg.click("a[data-tab=qui]"); pg.wait_for_selector("#quiBody table.qui", timeout=8000)
    B.ok(not errs, "aucune erreur JavaScript" + (f" : {errs[:3]}" if errs else ""))
    b.close()
raise SystemExit(0 if B.fin() else 1)
